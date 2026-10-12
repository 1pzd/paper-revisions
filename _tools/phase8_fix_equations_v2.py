#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Phase 8 Equation Recovery: Remove duplicates and reconstruct missing equations.

Current state (from diagnostic):
- P069-P070: DUPLICATE L_focal (7)
- P073-P074: DUPLICATE L_rank (8)
- P076 + P078: DUPLICATE L_total (9)
- Missing: equations (1), (2), (4), (5), (6)

Strategy:
1. Delete duplicate equation paragraphs
2. Identify which equations (1)-(6) are missing
3. Reconstruct missing equations in correct locations
4. Verify final count = 9 with numbers (1)-(9)
"""
from docx import Document
from docx.oxml.ns import qn
import sys
import io
import os
import re

sys.path.insert(0, os.path.dirname(__file__))
import omml

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

def get_equation_number(p):
    """Extract equation number from paragraph text"""
    text = ''.join(n.text or '' for n in p._p.iter() if n.tag in (W + 't', M + 't'))
    m = re.search(r'\((\d+)\)', text)
    return int(m.group(1)) if m else None

def has_omath(p):
    """Check if paragraph contains oMath nodes"""
    return bool(p._p.findall('.//' + qn('m:oMath')))

def delete_paragraph(p):
    """Delete a paragraph from document"""
    p._p.getparent().remove(p._p)

def main():
    doc = Document('../CLIP.docx')
    
    print("=== Phase 8 Equation Recovery ===\n")
    
    # Step 1: Identify all equation paragraphs
    print("1. Scanning for equation paragraphs...")
    eq_paras = []
    for i, p in enumerate(doc.paragraphs):
        if has_omath(p):
            eq_num = get_equation_number(p)
            eq_paras.append((i, p, eq_num))
            print(f"   P{i:03d}: equation ({eq_num if eq_num else '?'})")
    
    print(f"\nFound {len(eq_paras)} equation paragraphs")
    
    # Step 2: Remove duplicates (keep first occurrence of each number)
    print("\n2. Removing duplicate equations...")
    seen_numbers = set()
    to_delete = []
    
    for idx, p, eq_num in eq_paras:
        if eq_num is not None:
            if eq_num in seen_numbers:
                print(f"   DELETE P{idx:03d}: duplicate equation ({eq_num})")
                to_delete.append(p)
            else:
                seen_numbers.add(eq_num)
                print(f"   KEEP P{idx:03d}: equation ({eq_num})")
        else:
            # Unnumbered equation - keep for now
            print(f"   KEEP P{idx:03d}: unnumbered equation")
    
    for p in to_delete:
        delete_paragraph(p)
    
    print(f"Deleted {len(to_delete)} duplicate paragraphs")
    
    # Step 3: Identify missing equations
    print("\n3. Identifying missing equations...")
    missing = []
    for i in range(1, 10):
        if i not in seen_numbers:
            missing.append(i)
            print(f"   Missing: equation ({i})")
    
    if not missing:
        print("   All equations (1)-(9) present")
        doc.save('../CLIP.docx')
        print("\n=== Recovery complete (no missing equations) ===")
        return 0
    
    # Step 4: Reconstruct missing equations
    print(f"\n4. Reconstructing {len(missing)} missing equations...")
    
    # Reload document after deletions
    doc.save('../CLIP.docx')
    doc = Document('../CLIP.docx')
    
    # Find equation style reference
    eq_style_ref = None
    for p in doc.paragraphs:
        if has_omath(p):
            eq_style_ref = p
            break
    
    if eq_style_ref is None:
        print("[ERROR] No equation paragraph found for style reference")
        return 1
    
    L = '\u2112'  # Script L
    
    # Reconstruct each missing equation
    # Note: We need to find correct insertion points based on surrounding text
    
    # Helper function to insert equation
    def insert_eq_after_text(anchor_text, elements, number):
        for i, p in enumerate(doc.paragraphs):
            if anchor_text in p.text and not has_omath(p):
                # Insert equation after this paragraph
                eq_p = omml.equation_paragraph(doc, elements, number, pad_lead=47, pad_trail=12)
                # Copy style
                from docx.oxml import OxmlElement
                import copy
                src_pr = eq_style_ref._p.find(W + 'pPr')
                if src_pr is not None:
                    dst_pr = eq_p._p.find(W + 'pPr')
                    if dst_pr is None:
                        dst_pr = OxmlElement(W + 'pPr')
                        eq_p._p.insert(0, dst_pr)
                    s = src_pr.find(W + 'pStyle')
                    if s is not None:
                        old = dst_pr.find(W + 'pStyle')
                        if old is not None:
                            dst_pr.remove(old)
                        dst_pr.insert(0, copy.deepcopy(s))
                
                p._p.addnext(eq_p._p)
                print(f"   Inserted equation ({number}) after P{i:03d}")
                return True
        print(f"   [WARN] Anchor not found for equation ({number}): {anchor_text[:50]}")
        return False
    
    # Reconstruct missing equations based on standard locations
    if 1 in missing:
        # y = f(x; θ) ∈ {G0, ..., G5}
        eq1 = [omml.run('y = f(x; \u03b8) \u2208 {G'), omml.sub('0', ''),
               omml.run(', ..., G'), omml.sub('5', ''), omml.run('}')]
        insert_eq_after_text("Our objective is to predict", eq1, 1)
    
    if 2 in missing:
        # z_v = W_v · h_v
        eq2 = [omml.sub('z', 'v'), omml.run(' = '), omml.sub('W', 'v'),
               omml.run(' \u00b7 '), omml.sub('h', 'v')]
        insert_eq_after_text("The encoder", eq2, 2)
    
    if 4 in missing:
        # L_align = -log(...)
        eq4 = [omml.sub(L, 'align'), omml.run(' = \u2212log'),
               omml.frac([omml.run('exp('), omml.sub('z', 'v'), omml.run('\u00b7'),
                         omml.sub('z', 't'), omml.run('/\u03c4)')],
                        [omml.nary('\u2211', 'k', [omml.run('exp('), omml.sub('z', 'v'),
                                                   omml.run('\u00b7'), omml.sub('z', 'k'),
                                                   omml.run('/\u03c4)')])])]
        insert_eq_after_text("text descriptions are encoded", eq4, 4)
    
    if 5 in missing:
        # z_t^+ = λ·z_t + (1-λ)·c_y
        eq5 = [omml.sup([omml.sub('z', 't')], '+'), omml.run(' = \u03bb\u00b7'),
               omml.sub('z', 't'), omml.run(' + (1\u2212\u03bb)\u00b7'),
               omml.sub('c', 'y')]
        insert_eq_after_text("learnable scaling factor", eq5, 5)
    
    if 6 in missing:
        # c_y^(t) = m·c_y^(t-1) + (1-m)·z_t
        eq6 = [omml.sup([omml.sub('c', 'y')], '(t)'), omml.run(' = m\u00b7'),
               omml.sup([omml.sub('c', 'y')], '(t\u22121)'),
               omml.run(' + (1\u2212m)\u00b7'), omml.sub('z', 't')]
        insert_eq_after_text("momentum-updated global class centers", eq6, 6)
    
    # Save and verify
    doc.save('../CLIP.docx')
    print("\n=== Verifying final state ===")
    
    doc2 = Document('../CLIP.docx')
    final_eq_nums = []
    for p in doc2.paragraphs:
        if has_omath(p):
            eq_num = get_equation_number(p)
            if eq_num is not None:
                final_eq_nums.append(eq_num)
    
    final_eq_nums.sort()
    print(f"Final equation numbers: {final_eq_nums}")
    print(f"Paragraphs: {len(doc2.paragraphs)}")
    print(f"Tables: {len(doc2.tables)}")
    
    if final_eq_nums == list(range(1, 10)):
        print("\n[PASS] All equations (1)-(9) present and correctly numbered")
        return 0
    else:
        print(f"\n[FAIL] Expected [1, 2, 3, 4, 5, 6, 7, 8, 9], got {final_eq_nums}")
        return 1

if __name__ == '__main__':
    sys.exit(main())
