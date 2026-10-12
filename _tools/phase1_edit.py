# -*- coding: utf-8 -*-
"""Phase 1: insert L_focal / L_rank formulas, expand training details,
fix miscitations, add focal-loss reference, unify equation numbering.
Idempotent: re-running skips completed steps.
"""
import sys, io, re, shutil, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

sys.path.insert(0, r"E:\修改自己的论文\_tools")
import omml

SRC = r"E:\修改自己的论文\CLIP.docx"
BAK = r"E:\修改自己的论文\_backup\CLIP_original.docx"
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

# ---------- backup ----------
import os
os.makedirs(os.path.dirname(BAK), exist_ok=True)
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
    print(f"[backup] created {BAK}")
else:
    print(f"[backup] exists")

doc = Document(SRC)

def para_by_text(substr):
    for p in doc.paragraphs:
        if substr in p.text:
            return p
    return None

def all_text_nodes(p):
    out = []
    for node in p._p.iter():
        if node.tag in (W + 't', M + 't', W + 'instrText'):
            out.append(node)
    return out

def replace_in_runs(p, old, new, must=True):
    hits = 0
    for node in p._p.iter():
        if node.tag in (W + 't', M + 't') and node.text and old in node.text:
            node.text = node.text.replace(old, new)
            hits += 1
    print(f"  replace {old[:40]!r} -> {new[:40]!r}: {hits} hit(s)")
    if must and hits == 0:
        raise SystemExit(f"FAILED to find text: {old!r}")

def copy_pstyle(src_p, dst_p):
    src = src_p._p.find(W + 'pPr')
    if src is None:
        return
    s = src.find(W + 'pStyle')
    if s is None:
        return
    dpr = dst_p._p.find(W + 'pPr')
    if dpr is None:
        dpr = dst_p._p.makeelement(W + 'pPr', {})
        dst_p._p.insert(0, dpr)
    old = dpr.find(W + 'pStyle')
    if old is not None:
        dpr.remove(old)
    dpr.insert(0, copy.deepcopy(s))

def insert_paras_after(anchor, builders):
    """builders: list of callables(doc)->paragraph, inserted in order after anchor."""
    cursor = anchor._p
    made = []
    for b in builders:
        np = b(doc)
        cursor.addnext(np._p)
        cursor = np._p
        made.append(np)
    return made

# ============================================================
# STEP A: 3.4 insert L_focal (7) and L_rank (8) formulas
# ============================================================
print("\n[A] 3.4 rank-loss formulas")
if para_by_text("The bidirectional rank-aware term penalizes"):
    print("  already done, skip")
else:
    anchor = para_by_text("Figure 2 illustrates the physical simulation")
    assert anchor is not None, "anchor P064 not found"
    eq_ref = para_by_text("This scheme utilizes Focal Loss")  # style ref (p1a)
    eq_para_ref = None
    for p in doc.paragraphs:
        if p.style and p.style.name == 'equation':
            eq_para_ref = p
            break

    L = '\u2112'  # script L
    def mk(text):
        def _b(d):
            p = d.add_paragraph(text, style=eq_ref.style.name)
            copy_pstyle(eq_ref, p)
            return p
        return _b
    def mk_eq(elements, number):
        def _b(d):
            p = omml.equation_paragraph(d, elements, number, pad_lead=47, pad_trail=12)
            copy_pstyle(eq_para_ref, p)
            return p
        return _b

    lead1 = ("To make the optimization objective explicit, the focal component and the "
             "bidirectional rank-aware term are defined separately. The focal component reads:")
    eq7 = [omml.sub(L, 'focal'), omml.run(' = \u2212'),
           omml.sub('\u03b1', 't'),
           omml.sup([omml.run('(1\u2212'), omml.sub('p', 't'), omml.run(')')], '\u03b3'),
           omml.run(' log '), omml.sub('p', 't')]
    exp1 = ("where p_t = p(y|x) denotes the predicted probability of the ground-truth grade, "
            "\u03b1_t is the class-balancing weight, and \u03b3 = 2.0 down-weights well-classified "
            "samples so that hard tail-grade cases dominate the gradient[21].")
    lead2 = ("The bidirectional rank-aware term penalizes the expected ordinal distance between "
             "the predicted distribution and the ground-truth grade y:")
    eq8 = [omml.sub(L, 'rank'), omml.run(' = '),
           omml.nary('\u2211', 'c<y', [omml.run('(y\u2212c) p(c|x)')]),
           omml.run(' + '),
           omml.nary('\u2211', 'c>y', [omml.run('(c\u2212y) p(c|x)')])]
    exp2 = ("The first sum charges probability mass placed on grades below the true grade "
            "(underestimation), and the second charges mass placed above it (overestimation), so "
            "both directions are constrained simultaneously. The penalty grows linearly with the "
            "grade distance |c \u2212 y|: adjacent confusions are tolerated, whereas remote "
            "confusions are heavily penalized. Minimizing the rank term therefore concentrates "
            "the probability mass on y and shapes a unimodal distribution centered on the ground "
            "truth, eliminating logically inconsistent long-span errors (e.g., confusing Grade 0 "
            "with Grade 4). Equivalently, the term equals the expected ordinal distance "
            "E[|c \u2212 y|] under p(c|x), consistent with the objective stated in Section 3.1.")

    insert_paras_after(anchor, [mk(lead1), mk_eq(eq7, 7), mk(exp1),
                                mk(lead2), mk_eq(eq8, 8), mk(exp2)])
    print("  inserted 6 paragraphs (lead1, eq7, exp1, lead2, eq8, exp2)")

    # P066 cross references
    p66 = para_by_text("This scheme utilizes Focal Loss")
    replace_in_runs(p66, ") to mitigate class imbalance while leveraging the rank-aware term (",
                    ") (equation 7) to mitigate class imbalance while leveraging the rank-aware term (")
    replace_in_runs(p66, ") for ordinal constraints. With ",
                    ") (equation 8) for ordinal constraints. With ")
    replace_in_runs(p66, ", the system optimizes",
                    " (equation 9), the system optimizes")

# ============================================================
# STEP B: citation fixes
# ============================================================
print("\n[B] citation fixes")
p24 = para_by_text("Our objective is to predict")
if p24 is not None:
    hits = 0
    for node in p24._p.iter():
        if node.tag in (W + 't', M + 't') and node.text and '[14]' in node.text:
            node.text = node.text.replace('[14]', '[19]')
            hits += 1
    print(f"  P024 [14]->[19]: {hits} hit(s)")
    if hits == 0 and 'errors[19]' not in p24.text:
        # diagnostic: show run fragments near the citation
        frags = [n.text for n in p24._p.iter() if n.tag in (W + 't', M + 't') and n.text]
        print("  DIAG run fragments:", frags)
        raise SystemExit("FAILED: citation [14] fragment not located")
    p16 = para_by_text("VLMs like CLIP represent")
    if 'downstream adaptation[14]' not in p16.text:
        h2 = 0
        for node in p16._p.iter():
            if node.tag in (W + 't', M + 't') and node.text and 'downstream adaptation' in node.text:
                node.text = node.text.replace('downstream adaptation', 'downstream adaptation[14]')
                h2 += 1
        print(f"  P016 add [14]: {h2} hit(s)")
        if h2 == 0:
            frags = [n.text for n in p16._p.iter() if n.tag in (W + 't', M + 't') and n.text]
            print("  DIAG run fragments:", frags)
            raise SystemExit("FAILED: 'downstream adaptation' fragment not located")
else:
    print("  P024 not found")

# ============================================================
# STEP C: reference [21] focal loss
# ============================================================
print("\n[C] reference [21]")
if para_by_text("Focal loss for dense object detection") is None:
    last_ref = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("Zhang, H., Cisse"):
            last_ref = p
    assert last_ref is not None, "ref [20] anchor not found"
    newref = doc.add_paragraph(
        "Lin, T.-Y., Goyal, P., Girshick, R., He, K., Doll\u00e1r, P.: Focal loss for dense object "
        "detection. In: Proceedings of the IEEE International Conference on Computer Vision "
        "(ICCV), pp. 2980\u20132988 (2017)", style=last_ref.style.name)
    copy_pstyle(last_ref, newref)
    last_ref._p.addnext(newref._p)
    print("  added focal-loss reference")
else:
    print("  already done, skip")

# ============================================================
# STEP D: 4.2 training details expansion
# ============================================================
print("\n[D] 4.2 training details")
if para_by_text("All slit-lamp images were resized") is None:
    anchor = para_by_text("Experiments were conducted on an NVIDIA RTX 4050")
    assert anchor is not None
    detail = (
        "All slit-lamp images were resized to 224 \u00d7 224 before entering the encoders. The "
        "global branch partially fine-tunes the CLIP ViT-L/14 image encoder\u2014the first 12 "
        "transformer blocks are frozen while the remaining 12 blocks, the attention fusion layer, "
        "and the projection layers are trained\u2014and the text encoder provides a fixed "
        "pretrained semantic space\u3010\u5f85\u8865\uff1a\u8bf7\u786e\u8ba4\u6587\u672c\u7f16\u7801"
        "\u5668\u662f\u5426\u53c2\u4e0e\u8bad\u7ec3\u3011. The dataset was split into training, "
        "validation and test subsets\u3010\u5f85\u8865\uff1a\u8bf7\u8865\u5145\u5212\u5206\u6bd4"
        "\u4f8b\u3011; models were trained for \u3010\u5f85\u8865\uff1a\u8bf7\u8865\u5145\u8bad"
        "\u7ec3\u8f6e\u6570\u3011 epochs with weight decay \u3010\u5f85\u8865\uff1a\u8bf7\u8865"
        "\u5145\u6743\u91cd\u8870\u51cf\u7cfb\u6570\u3011, and the checkpoint with the best "
        "validation performance was used for evaluation. Random horizontal flipping and mild "
        "color jitter were applied as data augmentation\u3010\u5f85\u8865\uff1a\u8bf7\u786e\u8ba4"
        "\u589e\u5f3a\u7b56\u7565\u3011. The momentum coefficient of the prototype smoothing "
        "(global class center) update was set to 0.9\u3010\u5f85\u8865\uff1a\u8bf7\u786e\u8ba4"
        "\u52a8\u91cf\u7cfb\u6570\u3011. The temperature \u03c4 of the semantic projection was "
        "initialized to 0.07 and learned end-to-end; the loss weight \u03bb was 1.0; and K = 6 "
        "patches were retained per image. All results are averaged over \u3010\u5f85\u8865\uff1a"
        "\u8bf7\u8865\u5145\u91cd\u590d\u6b21\u6570\u4e0e\u968f\u673a\u79cd\u5b50\u3011 runs.")
    np = doc.add_paragraph(detail, style=anchor.style.name)
    copy_pstyle(anchor, np)
    anchor._p.addnext(np._p)
    print("  inserted training-details paragraph")
else:
    print("  already done, skip")

# ============================================================
# STEP E: unify equation numbers 1..N as plain text
# ============================================================
print("\n[E] equation numbering")
eq_paras = [p for p in doc.paragraphs if p.style and p.style.name == 'equation']
print(f"  found {len(eq_paras)} equation paragraphs")

def strip_field_runs(p):
    """Remove w:r elements that contain fldChar/instrText (real Word fields)."""
    removed = 0
    for r in list(p._p.iter(W + 'r')):
        if r.find(W + 'fldChar') is not None or r.find(W + 'instrText') is not None:
            r.getparent().remove(r)
            removed += 1
    return removed

for i, p in enumerate(eq_paras, start=1):
    removed = strip_field_runs(p)
    new_label = f'({i})'
    tail = None
    for node in list(p._p.iter()):
        if node.tag not in (W + 't', M + 't'):
            continue
        t = node.text or ''
        if 'SEQ' in t:
            node.text = re.sub(r'\(\s*SEQ.*?\d\)\s*$', lambda mm: ' ' * 6 + new_label, t,
                               flags=re.S)
            tail = node
        elif re.search(r'\(\d+\)\s*$', t):
            node.text = re.sub(r'\(\d+\)\s*$', new_label, t)
            tail = node
    if tail is None:
        # no trailing number found -> append one to last text node
        nodes = [n for n in p._p.iter() if n.tag in (W + 't', M + 't')]
        if nodes:
            nodes[-1].text = (nodes[-1].text or '') + ' ' * 6 + new_label
    print(f"  eq {new_label}: fields_removed={removed}")

doc.save(SRC)
print(f"\n[saved] {SRC}")

# ============================================================
# VERIFY
# ============================================================
doc2 = Document(SRC)
n_tab = len(doc2.tables)
n_img = len(doc2.inline_shapes)
n_para = len(doc2.paragraphs)
eq_nums = []
for p in doc2.paragraphs:
    if p.style and p.style.name == 'equation':
        t = ''.join(n.text or '' for n in p._p.iter() if n.tag in (W + 't', M + 't'))
        m = re.search(r'\((\d+)\)\s*$', t)
        eq_nums.append(m.group(1) if m else '??')
print(f"[verify] paragraphs={n_para} tables={n_tab} images={n_img}")
print(f"[verify] equation numbers: {eq_nums}")
ok = n_tab == 3 and n_img == 2 and eq_nums == [str(x) for x in range(1, len(eq_nums) + 1)]
print("[verify] PASS" if ok else "[verify] CHECK FAILED")
