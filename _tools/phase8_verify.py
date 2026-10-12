#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Phase 8 Verification: 检查 NC+P 联合严重程度术语修正完整性
"""
from docx import Document
from docx.oxml.ns import qn
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def check_equations(doc):
    """验证公式 (1)-(9) 完好"""
    eq_count = 0
    for p in doc.paragraphs:
        for run in p.runs:
            if run.element.findall('.//' + qn('m:oMath')):
                eq_count += 1
    return eq_count

def main():
    doc = Document('../CLIP.docx')
    
    print("=== Phase 8 Verification ===\n")
    
    checks = []
    
    # 1. 检查 'nuclear grade' 残留（排除引用上下文）
    print("1. Checking for residual 'nuclear grade' terminology...")
    nuclear_count = 0
    for i, p in enumerate(doc.paragraphs):
        if "nuclear grade" in p.text.lower():
            # 允许在参考文献或引用中出现
            if not (p.text.startswith("[") or "et al" in p.text or "References" in p.text):
                print(f"   [WARN] P{i:03d}: {p.text[:80]}")
                nuclear_count += 1
    if nuclear_count == 0:
        print("   [OK] No residual 'nuclear grade' outside citations")
        checks.append(True)
    else:
        print(f"   [FAIL] Found {nuclear_count} residual occurrences")
        checks.append(False)
    
    # 2. 验证 Table 1 prompts 已更新为 G0-G5
    print("\n2. Checking Table 1 prompts...")
    if len(doc.tables) >= 1:
        tbl = doc.tables[0]
        g_count = 0
        for i in range(min(6, len(tbl.rows))):
            cell_text = tbl.cell(i, 1).text
            if f"G{i}" in cell_text and "severity" in cell_text:
                g_count += 1
        if g_count == 6:
            print(f"   [OK] All 6 prompts use G0-G5 severity terminology")
            checks.append(True)
        else:
            print(f"   [FAIL] Only {g_count}/6 prompts updated")
            checks.append(False)
    else:
        print("   [ERROR] Table 1 not found")
        checks.append(False)
    
    # 3. 验证 Table 5 caption
    print("\n3. Checking Table 5 caption...")
    caption_ok = False
    for p in doc.paragraphs:
        if "Confusion matrix" in p.text:
            if "joint severity" in p.text or ("weighted sum" in p.text and "NC" in p.text and "P" in p.text):
                print("   [OK] Table 5 caption updated with joint severity")
                caption_ok = True
                break
            elif "nuclear grades NC1" in p.text:
                print("   [FAIL] Table 5 still mentions 'nuclear grades NC1-NC6'")
                break
    checks.append(caption_ok)
    
    # 4. 验证公式 (1)-(9) 完好
    print("\n4. Checking equations...")
    eq_count = check_equations(doc)
    if eq_count == 9:
        print(f"   [OK] All 9 equations intact")
        checks.append(True)
    else:
        print(f"   [WARN] Found {eq_count} equations (expected 9)")
        checks.append(eq_count >= 9)
    
    # 5. 验证表格数量
    print("\n5. Checking table count...")
    if len(doc.tables) == 5:
        print("   [OK] 5 tables present")
        checks.append(True)
    else:
        print(f"   [WARN] Found {len(doc.tables)} tables (expected 5)")
        checks.append(False)
    
    # 6. 验证段落数（应该稳定在 134 左右）
    print("\n6. Checking paragraph count...")
    para_count = len(doc.paragraphs)
    if 130 <= para_count <= 140:
        print(f"   [OK] {para_count} paragraphs (expected ~134)")
        checks.append(True)
    else:
        print(f"   [WARN] {para_count} paragraphs (expected 130-140 range)")
        checks.append(False)
    
    # 7. 检查 P024 问题定义是否提到 NC + P
    print("\n7. Checking problem definition...")
    prob_def_ok = False
    for i, p in enumerate(doc.paragraphs):
        if "weighted sum" in p.text and "NC" in p.text and "P" in p.text and "severity" in p.text:
            print(f"   [OK] P{i:03d} contains joint severity definition")
            prob_def_ok = True
            break
    checks.append(prob_def_ok)
    
    # 总结
    print("\n" + "="*50)
    passed = sum(checks)
    total = len(checks)
    if passed == total:
        print(f"[PASS] Phase 8 verification complete: {passed}/{total} checks passed")
        return 0
    else:
        print(f"[PARTIAL] Phase 8 verification: {passed}/{total} checks passed")
        return 1

if __name__ == '__main__':
    sys.exit(main())
