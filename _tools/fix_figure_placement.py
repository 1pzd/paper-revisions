#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Fix figure insertion issues:
1. Move Fig.4 to correct location (after dataset description in 4.1)
2. Insert Fig.5 after Table 5
3. Remove duplicate Fig.2 description paragraph
"""
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
import io
import os

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

def find_para_containing(doc, anchor_text, start_idx=0):
    """Find paragraph containing anchor text"""
    for i in range(start_idx, len(doc.paragraphs)):
        if anchor_text.lower() in doc.paragraphs[i].text.lower():
            return i, doc.paragraphs[i]
    return None, None

def find_table_index(doc, caption_keyword):
    """Find table by searching for caption before it"""
    for i, p in enumerate(doc.paragraphs):
        if caption_keyword.lower() in p.text.lower() and "table" in p.text.lower():
            # The table should be nearby (within next 3 paragraphs)
            # Return the paragraph index where we should insert after
            return i
    return None

def insert_figure_after(doc, para_idx, image_path, caption_text, width_cm=12):
    """Insert figure with caption after specified paragraph index"""
    anchor_para = doc.paragraphs[para_idx]
    
    # Create figure paragraph
    fig_para = doc.add_paragraph()
    anchor_para._p.addnext(fig_para._p)
    
    # Add image
    run = fig_para.add_run()
    run.add_picture(image_path, width=Cm(width_cm))
    fig_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Create caption paragraph
    caption_para = doc.add_paragraph()
    fig_para._p.addnext(caption_para._p)
    caption_para.add_run(caption_text)
    caption_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    return fig_para, caption_para

def delete_paragraph(para):
    """Delete a paragraph from document"""
    p_element = para._p
    p_element.getparent().remove(p_element)

def main():
    doc = Document('../CLIP.docx')
    
    print("=== Fixing Figure Insertion Issues ===\n")
    
    base_path = r'E:\修改自己的论文\data_in'
    
    # Step 1: Remove duplicate "Figure 2 illustrates" paragraph at P066
    print("1. Removing duplicate Figure 2 description...")
    idx, p = find_para_containing(doc, "Figure 2 illustrates the physical simulation")
    if p and "Figure 2 illustrates" in p.text:
        delete_paragraph(p)
        print(f"   Deleted duplicate description at P{idx:03d}")
    
    # Step 2: Find and move Fig.4 to correct location
    print("\n2. Moving Fig.4 to correct location...")
    
    # Find current Fig.4 caption
    idx_fig4, p_fig4 = find_para_containing(doc, "Fig. 4. Long-tailed distribution")
    if p_fig4:
        # Find the image paragraph before this caption (should be right before)
        fig4_img_para = None
        for i in range(max(0, idx_fig4-2), idx_fig4):
            para = doc.paragraphs[i]
            # Check if paragraph contains image
            if para._element.xpath('.//pic:pic'):
                fig4_img_para = para
                break
        
        if fig4_img_para:
            # Delete current Fig.4 and its caption
            delete_paragraph(fig4_img_para)
            delete_paragraph(p_fig4)
            print(f"   Removed misplaced Fig.4 from P{idx_fig4:03d} area")
            
            # Re-insert at correct location (after dataset description in 4.1)
            # Look for dataset description with distribution numbers
            idx_target, _ = find_para_containing(doc, "230, 759, 548, 101, 20, 15")
            if not idx_target:
                idx_target, _ = find_para_containing(doc, "the six severity grades")
            
            if idx_target:
                caption = "Fig. 4. Long-tailed distribution of the 1,673-image dataset over the six severity grades: the two head grades cover 78.1% of samples, while the two most severe grades cover only 2.1%."
                insert_figure_after(doc, idx_target, os.path.join(base_path, 'fig4.jpg'), caption, width_cm=10)
                print(f"   Re-inserted Fig.4 after P{idx_target:03d} (dataset description)")
            else:
                print("   [WARN] Could not find dataset description anchor")
    
    # Step 3: Insert Fig.5 after Table 5
    print("\n3. Inserting Fig.5 after Table 5...")
    
    # Find Table 5 - it should be after "Confusion matrix" caption
    table5_idx = find_table_index(doc, "confusion matrix")
    
    if table5_idx:
        # Table 5 is a 7x7 table, so skip ahead past it
        # Tables appear as inline elements, so we need to find the paragraph after the table
        insert_after_idx = table5_idx + 1
        
        # Skip ahead to find the first paragraph after the table content
        # Usually there are a few blank or short paragraphs after a table
        for i in range(table5_idx + 1, min(table5_idx + 10, len(doc.paragraphs))):
            # Look for a paragraph that's likely after the table
            # (empty, or starts a new section, or has substantial text)
            para_text = doc.paragraphs[i].text.strip()
            if len(para_text) > 100 or para_text == "":
                insert_after_idx = i
                break
        
        caption = "Fig. 5. Confusion matrix on the held-out test set (335 images): errors concentrate on adjacent grades (40 of 44), and no sample deviates by three or more grades."
        insert_figure_after(doc, insert_after_idx, os.path.join(base_path, 'fig5.png'), caption, width_cm=10)
        print(f"   Inserted Fig.5 after P{insert_after_idx:03d} (Table 5 area)")
    else:
        print("   [WARN] Could not find Table 5 location")
    
    # Save
    doc.save('../CLIP.docx')
    print("\n=== Changes saved ===")
    
    # Verify
    print("\n=== Verification ===")
    doc2 = Document('../CLIP.docx')
    
    fig_captions = []
    for i, p in enumerate(doc2.paragraphs):
        if p.text.startswith("Fig. "):
            fig_num = p.text.split('.')[0].split()[-1]
            fig_captions.append((i, fig_num, p.text[:60]))
    
    print(f"Figure captions: {len(fig_captions)}")
    for idx, num, text in fig_captions:
        print(f"  P{idx:03d}: Fig.{num} - {text}")
    
    print(f"\nInline images: {len(doc2.inline_shapes)}")
    print(f"Paragraphs: {len(doc2.paragraphs)}")
    
    if len(fig_captions) == 5 and len(doc2.inline_shapes) == 5:
        print("\n[PASS] All 5 figures correctly placed")
        return 0
    else:
        print(f"\n[WARN] Expected 5 figures, found {len(fig_captions)} captions and {len(doc2.inline_shapes)} images")
        return 1

if __name__ == '__main__':
    sys.exit(main())
