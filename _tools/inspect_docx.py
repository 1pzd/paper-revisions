# -*- coding: utf-8 -*-
"""Inspect CLIP.docx structure: paragraphs, OMML equations, tables, images."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

DOC = r"E:\修改自己的论文\CLIP.docx"
doc = Document(DOC)

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

def para_info(p, idx):
    texts = []
    math_count = 0
    math_texts = []
    for node in p._p.iter():
        tag = node.tag
        if tag == W + 't':
            if node.text:
                texts.append(node.text)
        elif tag == M + 't':
            math_count += 1
            if node.text:
                math_texts.append(node.text)
    full = ''.join(texts)
    math_str = ''.join(math_texts)
    style = p.style.name if p.style else '?'
    return idx, style, full, math_count, math_str

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    idx, style, full, mc, ms = para_info(p, i)
    line = f"[{i:03d}]({style}) {full[:110]}"
    if mc:
        line += f" <<MATH x{mc}: {ms[:60]}>>"
    # detect images
    blips = p._p.findall('.//' + '{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    if blips:
        line += f" <<IMG x{len(blips)}>>"
    print(line)

print("\n=== TABLES ===")
for ti, t in enumerate(doc.tables):
    print(f"--- Table {ti}: {len(t.rows)} rows x {len(t.columns)} cols ---")
    for r in t.rows[:3]:
        cells = [c.text.strip().replace('\n', ' ')[:22] for c in r.cells]
        print("   | " + " | ".join(cells))

print("\n=== INLINE SHAPES (images) ===")
for si, s in enumerate(doc.inline_shapes):
    print(f"shape {si}: type={s.type}, w={s.width}, h={s.height}")
