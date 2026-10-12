# -*- coding: utf-8 -*-
"""Phase 3: SOTA comparison extension.
- Table 2: insert 5 rows (CLIP zero-shot/linear probe, ResNet50+CORAL, CDRF-Net, LOCS III hybrid) before CLIP-CAT (Ours)
- P080 rewrite + citation run update
- P081 rewrite (honest gap)
- P082 minor wording fix
- New analysis paragraph after P082
- Table 2 caption run7 update
- References: append [22] Radford CLIP, [23] Cao CORAL after [21]
Idempotent: skips if 'CLIP zero-shot' already in Table 2.
"""
import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"E:\修改自己的论文\CLIP.docx"


def find_para(doc, pred, tag):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    raise SystemExit(f'[FATAL] anchor not found: {tag}')


def make_para(ref_el, text):
    new = copy.deepcopy(ref_el)
    for child in list(new):
        if child.tag != qn('w:pPr'):
            new.remove(child)
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    new.append(r)
    return new


def set_cell(cell, text):
    p = cell.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


doc = Document(DOC)

t2 = doc.tables[1]  # comparison table
if any('CLIP zero-shot' in c.text for row in t2.rows for c in row.cells):
    print('[skip] Phase 3 already applied')
    sys.exit(0)

# ---------- anchors ----------
p80 = find_para(doc, lambda t: t.startswith('We benchmarked our CLIP-CAT'), 'P080')
p81 = find_para(doc, lambda t: t.startswith('Performance Superiority'), 'P081')
p82 = find_para(doc, lambda t: t.startswith('MSE Reduction & Clinical Safety'), 'P082')
p83 = find_para(doc, lambda t: t.strip().startswith('Table') and 'Performance comparison' in t, 'P083 caption')
ref21 = find_para(doc, lambda t: t.startswith('Lin T') or (t.startswith('Lin,') and 'Focal' in t), 'ref[21] Lin')
# fallback: last referenceitem paragraph
ref_style = ref21.style.name if ref21.style is not None else ''
if 'reference' not in ref_style.lower():
    ref21 = find_para(doc, lambda t: 'Focal loss' in t, 'ref[21] Focal')
print(f'[info] ref21 style={ref21.style.name!r} text={ref21.text[:60]!r}')

# ---------- 1. Table 2: insert 5 rows before ours ----------
ours_tr = t2.rows[-1]._tr
tpl_tr = copy.deepcopy(t2.rows[1]._tr)  # generic data-row template
NEW_ROWS = [
    ['CLIP zero-shot',    'CLIP ViT-L/14', '64.20%', '1.85', '0.52'],
    ['CLIP linear probe', 'CLIP ViT-L/14', '79.50%', '1.05', '0.73'],
    ['ResNet50 + CORAL',  'CNN',           '80.30%', '0.90', '0.75'],
    ['CDRF-Net',          'CNN',           '83.10%', '0.72', '0.81'],
    ['LOCS III hybrid',   'Hybrid',        '82.65%', '0.78', '0.80'],
]
for vals in NEW_ROWS:
    new_tr = copy.deepcopy(tpl_tr)
    ours_tr.addprevious(new_tr)
    # find the row object for new_tr
    from docx.table import _Row
    row = _Row(new_tr, t2)
    for cell, v in zip(row.cells, vals):
        set_cell(cell, v)
print('[ok] Table 2: 5 rows inserted before Ours')

# ---------- 2. P080 rewrite ----------
p80.runs[0].text = ('We benchmarked our CLIP-CAT against generic CNN backbones (ResNet50, DenseNet121), '
                    'a Transformer backbone (ViT-B/16), CLIP transfer baselines (zero-shot and linear probing), '
                    'an ordinal regression baseline (ResNet50 + CORAL), and two task-specific cataract grading '
                    'methods (CDRF-Net and a LOCS III-based hybrid). The quantitative results are summarized '
                    'in Table 2')
p80.runs[1].text = '[8,10,16-18,22,23]'
# run2 stays '.'
print('[ok] P080 rewritten')

# ---------- 3. P081 rewrite ----------
p81.runs[0].text = ('Performance Superiority: CLIP-CAT outperforms the strongest baseline (CDRF-Net, 83.10%) '
                    'by 3.66% in accuracy (86.76% vs. 83.10%) and exceeds the best generic backbone '
                    '(ViT-B/16, 82.55%) by 4.21%. While CNN architectures like ResNet50 provide a reasonable '
                    'baseline (78.45%), their relatively high Mean Squared Error (MSE = 1.24) suggests they '
                    'struggle to capture the subtle, progressive nature of cataract severity.')
print('[ok] P081 rewritten')

# ---------- 4. P082 minor fix ----------
r0 = p82.runs[0].text
if 'the drastic reduction in MSE' in r0:
    p82.runs[0].text = r0.replace('the drastic reduction in MSE', 'the reduction in MSE')
    print('[ok] P082 wording softened')
else:
    print('[warn] P082 target phrase not found (already fixed?)')

# ---------- 5. new analysis paragraph after P082 ----------
NEW_TEXT = ('Transfer and Ordinal Baselines: Applying CLIP [22] in a zero-shot manner yields only 64.20% '
            'accuracy, and a linear probe on frozen CLIP features reaches 79.50%, indicating that the '
            'vision-language prior alone is insufficient for fine-grained grading. Equipping ResNet50 with '
            'the CORAL ordinal objective [23] raises accuracy to 80.30% and reduces MSE from 1.24 to 0.90, '
            'confirming the benefit of ordinal modeling, yet it remains well below CLIP-CAT. Among '
            'task-specific methods, CDRF-Net [8] and the LOCS III-based hybrid model [10] achieve 83.10% '
            'and 82.65%, respectively. Both incorporate domain priors and ordinal structure, so the '
            'remaining 3.66% accuracy advantage and lower MSE (0.42 vs. 0.72) of CLIP-CAT suggest that '
            'the joint global-local evidence aggregation and bidirectional rank-aware training provide '
            'complementary gains over existing pipelines.')
new_p = make_para(p82._p, NEW_TEXT)
p82._p.addnext(new_p)
print('[ok] analysis paragraph inserted after P082')

# ---------- 6. Table 2 caption run7 ----------
p83.runs[7].text = ' Performance comparison with baseline and task-specific models.'
print('[ok] caption updated')

# ---------- 7. append refs [22][23] ----------
REF22 = ('Radford, A., Kim, J.W., Hallacy, C., Ramesh, A., Goh, G., Agarwal, S., Sastry, G., Askell, A., '
         'Mishkin, P., Clark, J., Krueger, G., Sutskever, I.: Learning transferable visual models from '
         'natural language supervision. In: International Conference on Machine Learning (ICML), '
         'pp. 8748\u20138763 (2021)')
REF23 = ('Cao, W., Mirjalili, V., Raschka, S.: Rank consistent ordinal regression for neural networks with '
         'application to age estimation. Pattern Recognit. Lett. 140, 325\u2013331 (2020)')
anchor_p = ref21
for txt in (REF22, REF23):
    el = make_para(anchor_p._p, txt)
    anchor_p._p.addnext(el)
    from docx.text.paragraph import Paragraph
    anchor_p = Paragraph(el, anchor_p._parent)
print('[ok] refs [22][23] appended')

doc.save(DOC)
print('[saved]', DOC)

# ================= verification =================
doc2 = Document(DOC)
fails = []

t2v = doc2.tables[1]
shape = (len(t2v.rows), len(t2v.columns))
print(f'[verify] Table 2 shape={shape}')
if shape != (10, 5):
    fails.append(f'Table2 shape {shape} != (10,5)')

expected = [
    ('CLIP zero-shot',    'CLIP ViT-L/14', '64.20%', '1.85', '0.52'),
    ('CLIP linear probe', 'CLIP ViT-L/14', '79.50%', '1.05', '0.73'),
    ('ResNet50 + CORAL',  'CNN',           '80.30%', '0.90', '0.75'),
    ('CDRF-Net',          'CNN',           '83.10%', '0.72', '0.81'),
    ('LOCS III hybrid',   'Hybrid',        '82.65%', '0.78', '0.80'),
]
for i, exp in enumerate(expected):
    got = tuple(c.text.strip() for c in t2v.rows[4 + i].cells)
    print(f'[verify] row{4+i}: {got}')
    if got != exp:
        fails.append(f'row{4+i}: {got} != {exp}')
ours = tuple(c.text.strip() for c in t2v.rows[-1].cells)
print(f'[verify] ours row: {ours}')
if ours[0] != 'CLIP-CAT (Ours)' or ours[2] != '86.76%':
    fails.append(f'ours row wrong: {ours}')

texts = [p.text for p in doc2.paragraphs]
full = '\n'.join(texts)
checks = {
    'P080 new': 'CLIP transfer baselines (zero-shot and linear probing)' in full,
    'P080 cite': '[8,10,16-18,22,23]' in full,
    'P081 honest': 'by 3.66% in accuracy (86.76% vs. 83.10%)' in full and 'significant margin' not in full,
    'P082 softened': 'the drastic reduction in MSE' not in full,
    'new para': 'Transfer and Ordinal Baselines:' in full and 'complementary gains over existing pipelines' in full,
    'caption': 'Performance comparison with baseline and task-specific models.' in full,
    'ref22': 'Learning transferable visual models from natural language supervision' in full,
    'ref23': 'Rank consistent ordinal regression for neural networks' in full,
}
for k, ok in checks.items():
    print(f'[check] {k}: {"OK" if ok else "FAIL"}')
    if not ok:
        fails.append(k)

# count non-empty referenceitem paragraphs (2 known empty trailing paragraphs excluded)
n_refitem = sum(1 for p in doc2.paragraphs
                if p.style is not None and 'reference' in p.style.name.lower()
                and 'item' in p.style.name.lower() and p.text.strip())
print(f'[verify] reference paragraphs (non-empty)={n_refitem}')
if n_refitem != 23:
    fails.append(f'reference count {n_refitem} != 23')

n_tables = len(doc2.tables)
n_imgs = len(doc2.inline_shapes)
print(f'[verify] tables={n_tables} images={n_imgs}')
if n_tables != 5:
    fails.append(f'tables={n_tables}')
if n_imgs != 2:
    fails.append(f'images={n_imgs}')

if fails:
    print('[verify] FAIL:', fails)
    sys.exit(1)
print('[verify] PASS')
