import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

for i in (75, 77, 78, 100):
    p = doc.paragraphs[i]
    print(f'===== P{i} [{p.style.name}] =====')
    print(p.text)
    print()

print('===== keyword scan =====')
keys = ['early stop', 'weight decay', 'dropout', 'Mixup', 'augment', 'flip', 'jitter',
        'cross-validation', 'fold', 'external', 'multi-center', 'overfit', 'regulariz',
        'Limitation', 'frozen', 'freeze', 'trainable']
for k in keys:
    hits = [i for i, p in enumerate(doc.paragraphs) if k.lower() in p.text.lower()]
    print(f'{k!r}: {hits}')
