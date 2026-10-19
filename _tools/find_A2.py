# -*- coding: utf-8 -*-
"""Find the class-weights / G-distribution sentences in CLIP.docx (plain text + OMML)."""
from docx import Document

KEYS = ['focal', 'sampl', 'imbalanc', 'w_0', '27.9', '37.2', 'distribution of the',
        'long-tail', 'long tail', 'class weight', 'weighted', 'n_g', 'severity grades is']

def has_math(p):
    return b'm:oMath' in p._p.xml.encode('utf-8', 'ignore')

doc = Document('../CLIP.docx')
print(f'paragraphs={len(doc.paragraphs)}')
for i, p in enumerate(doc.paragraphs):
    low = p.text.lower()
    if any(k in low for k in KEYS):
        print(f'\n[P{i:03d}] math={has_math(p)} runs={len(p.runs)}')
        print('  FULL:', p.text)
        for j, r in enumerate(p.runs):
            flags = ('B' if r.bold else '-') + ('I' if r.italic else '-') + ('S' if r.font.subscript else '-')
            print(f'  R{j:02d} [{flags}] {r.text!r}')
print('\n=== done ===')
