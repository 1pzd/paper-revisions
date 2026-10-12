# -*- coding: utf-8 -*-
"""Fix eq(1)/eq(3) number tails: blank literal SEQ garbage nodes. Then verify all."""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

SRC = r"E:\修改自己的论文\CLIP.docx"
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

doc = Document(SRC)
eq_paras = [p for p in doc.paragraphs if p.style and p.style.name == 'equation']

fixed = 0
for i, p in enumerate(eq_paras, start=1):
    nodes = [n for n in p._p.iter() if n.tag in (W + 't', M + 't')]
    joined = ''.join(n.text or '' for n in nodes)
    if re.search(r'\(\d+\)\s*$', joined):
        continue
    seq_idx = None
    for j, n in enumerate(nodes):
        if n.text and 'SEQ' in n.text:
            seq_idx = j
            n.text = ''
            break
    if seq_idx is None:
        print(f"[warn] eq {i}: no SEQ node; tail={joined[-40:]!r}")
        continue
    for n2 in nodes[seq_idx + 1:]:
        if n2.text and re.fullmatch(r'\d+', n2.text.strip()):
            n2.text = str(i)
            break
    fixed += 1
    print(f"[fix] eq {i}: SEQ garbage cleared, label=({i})")

doc.save(SRC)
print(f"[saved] fixed={fixed}")

# ---- full verify ----
doc2 = Document(SRC)
eq_nums = []
for p in doc2.paragraphs:
    if p.style and p.style.name == 'equation':
        t = ''.join(n.text or '' for n in p._p.iter() if n.tag in (W + 't', M + 't'))
        m = re.search(r'\((\d+)\)\s*$', t)
        eq_nums.append(m.group(1) if m else '??')
print(f"[verify] paragraphs={len(doc2.paragraphs)} tables={len(doc2.tables)} images={len(doc2.inline_shapes)}")
print(f"[verify] eq numbers: {eq_nums}")
ok = (len(doc2.tables) == 5 and len(doc2.inline_shapes) == 2
      and eq_nums == [str(x) for x in range(1, len(eq_nums) + 1)])
print("[verify] PASS" if ok else "[verify] FAIL")

# sanity: key texts present
whole = '\n'.join(p.text for p in doc2.paragraphs)
checks = [
    ("L_focal formula", 'focal' in whole and 'log' in whole),
    ("rank term text", "The bidirectional rank-aware term penalizes" in whole),
    ("eq7 ref", "(equation 7)" in whole),
    ("eq9 ref", "(equation 9)" in whole),
    ("citation [19]", "errors[19]." in whole or "errors[19]" in whole),
    ("citation [14] moved", "downstream adaptation[14]" in whole),
    ("ref [21]", "Focal loss for dense object detection" in whole),
    ("training details", "resized to 224" in whole),
    ("no leftover SEQ", "SEQ \"equation\"" not in whole),
]
for name, okc in checks:
    print(f"[check] {name}: {'OK' if okc else 'MISSING'}")
