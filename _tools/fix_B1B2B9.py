#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
B1 Ethics 升格 + B2 新增独立小节 + B9 关键词微调（docx 同步）
- B1: 数据集段后加伦理声明（委员会名称+批件号2026-ZFKY-008+赫尔辛基宣言+书面知情同意）
- B2: References 前新增 Author Contributions / Conflict of Interest / Funding / Data Availability；作者区加 ORCID
- B9: 关键词 3->6 并微调（limited-scale learning -> limited data learning，新增 class imbalance）
"""
from docx import Document
import sys, io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')


def find_para(doc, anchor_text, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start=start):
        if anchor_text in p.text:
            return i, p
    return None, None


def replace_in_para(para, old, new):
    if old not in para.text:
        return False
    new_text = para.text.replace(old, new)
    for run in para.runs:
        run.text = ''
    if para.runs:
        para.runs[0].text = new_text
    else:
        para.add_run(new_text)
    return True


def main():
    doc = Document('CLIP.docx')
    print('=== fix B1/B2/B9 on CLIP.docx ===\n')

    # ---------- B9: keywords ----------
    idx, p = find_para(doc, 'Keywords:')
    if p:
        ok = replace_in_para(
            p,
            'Keywords: Cataract Grading, Vision-Language Models, Limited-Scale Learning',
            'Keywords: Cataract Grading, Vision-Language Models, Limited Data Learning, '
            'Ordinal Regression, Class Imbalance, Medical Image Analysis',
        )
        print(f'[{"OK" if ok else "WARN"}] P{idx:03d} keywords updated (3 -> 6, wording tuned)')
    else:
        print('[WARN] keywords paragraph not found')

    # ---------- B2-ORCID: after address ----------
    idx, p = find_para(doc, 'School of Computer Science and Engineering')
    if p:
        nxt = doc.paragraphs[idx + 1]
        nxt.insert_paragraph_before(
            'ORCID: Zhida Pang, https://orcid.org/0009-0000-0788-1232', style=p.style
        )
        print(f'[OK] ORCID line inserted after P{idx:03d} (address)')
    else:
        print('[WARN] address paragraph not found')

    # ---------- B1: ethics statement after dataset paragraph ----------
    idx, p = find_para(doc, 'To validate our framework, we curated a private clinical dataset')
    if p:
        nxt = doc.paragraphs[idx + 1]
        nxt.insert_paragraph_before(
            'This study used a private, non-public dataset of cataract case photographs '
            'provided by a participating hospital. The study protocol was approved by the '
            "Medical Ethics Committee of Shenyang Fourth People's Hospital (approval no. "
            '2026-ZFKY-008) and conducted in accordance with the Declaration of Helsinki '
            '(as revised in 2013). Written informed consent for the use of the clinical '
            'images in research was obtained from the participating patients or their '
            'legally authorized representatives.',
            style=p.style,
        )
        print(f'[OK] ethics statement inserted after P{idx:03d} (dataset para)')
    else:
        print('[WARN] dataset paragraph not found')

    # ---------- B2: declarations before References ----------
    idx, p = find_para(doc, 'References')
    # 找到 heading1 的 References 节标题（排除正文里出现的词）
    if p is None or p.style.name != 'heading1':
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip() == 'References' and para.style.name == 'heading1':
                idx, p = i, para
                break
    if p is not None and p.style.name == 'heading1':
        body_style = doc.paragraphs[111].style  # conclusion body style (p1a)
        blocks = [
            ('Author Contributions', 'heading1'),
            ('Zhida Pang: methodology, software, investigation, formal analysis, '
             'writing – original draft. Shuang Wang: conceptualization, supervision, '
             'writing – review and editing. All authors read and approved the final '
             'manuscript.', body_style.name),
            ('Conflict of Interest', 'heading1'),
            ('The authors declare that the research was conducted in the absence of any '
             'commercial or financial relationships that could be construed as a '
             'potential conflict of interest.', body_style.name),
            ('Funding', 'heading1'),
            ('This research received no specific grant from any funding agency in the '
             'public, commercial, or not-for-profit sectors.', body_style.name),
            ('Data Availability', 'heading1'),
            ('The clinical image dataset used in this study is not publicly available '
             'because it contains confidential patient information and is subject to '
             'institutional restrictions; requests for access may be considered by the '
             'corresponding authors upon reasonable request and with approval of the '
             'participating hospital. The images were handled as confidential clinical '
             'data throughout the study. The source code used in this study is available '
             'at https://github.com/1pzd/-Project.git.', body_style.name),
        ]
        for text, style in blocks:
            p.insert_paragraph_before(text, style=style)
        print(f'[OK] 4 declaration sections inserted before P{idx:03d} References '
              f'({len(blocks)} paragraphs)')
    else:
        print('[WARN] References heading not found')

    doc.save('CLIP.docx')
    print('\n[saved] CLIP.docx')

    # ---- verify ----
    doc2 = Document('CLIP.docx')
    keys = ['Keywords', 'ORCID', '2026-ZFKY-008', 'Author Contributions',
            'Conflict of Interest', 'Funding', 'Data Availability',
            'github.com/1pzd']
    print('\n--- verify ---')
    for k in keys:
        hits = [f'P{i:03d}' for i, para in enumerate(doc2.paragraphs) if k in para.text]
        print(f'{k}: {hits if hits else "MISSING!"}')


if __name__ == '__main__':
    main()
