import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)
ps = doc.paragraphs

def show(i, full=True, runs=False):
    p = ps[i]
    t = p.text
    print(f'--- P{i:03d} [{p.style.name}] runs={len(p.runs)} ---')
    print(t if full else t[:200])
    if runs:
        for j, r in enumerate(p.runs):
            print(f'   run{j}: {r.text!r}')
    print()

# full numeric-narrative paragraphs
for i in (3, 77, 80, 97, 99, 101, 103, 107):
    show(i, full=True, runs=True)

# short confirmations
for i in (100, 24):
    show(i, full=True, runs=False)

print('=== Table1 r0 ===')
print(repr(doc.tables[0].rows[0].cells[0].text), '|', repr(doc.tables[0].rows[0].cells[1].text))
print('=== Table4 ===')
for r in doc.tables[3].rows:
    print(' | '.join(c.text.strip() for c in r.cells))
print('=== Table5 ===')
for r in doc.tables[4].rows:
    print(' | '.join(c.text.strip() for c in r.cells))

print('=== paragraphs containing 2,000 / 2000 / 400 / 53 / 84.9 / 15.1 / 70% / 1400 ===')
import re
pat = re.compile(r'2,000|2000|400|53 |84\.9|15\.1|70%|1400')
for i, p in enumerate(ps):
    if pat.search(p.text):
        print(f'P{i:03d}: {p.text[:150]}')
