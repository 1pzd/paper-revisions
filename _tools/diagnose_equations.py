#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Diagnostic: Check what's actually in the document after equation injection attempts.
"""
from docx import Document
from docx.oxml.ns import qn
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

doc = Document('../CLIP.docx')

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}'

print("=== Document Structure Diagnostic ===\n")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check for ANY math-related XML nodes
print("\n--- Searching for math-related XML nodes ---")
math_nodes = []
for i, p in enumerate(doc.paragraphs):
    # Check for oMath nodes
    omath = p._p.findall('.//' + qn('m:oMath'))
    omath_para = p._p.findall('.//' + qn('m:oMathPara'))
    
    if omath or omath_para:
        print(f"P{i:03d}: oMath={len(omath)}, oMathPara={len(omath_para)}")
        math_nodes.append(i)
        
        # Show paragraph text
        text = ''.join(n.text or '' for n in p._p.iter() if n.tag in (W + 't', M + 't'))
        print(f"       Text: {text[:80]}")

if not math_nodes:
    print("[NONE FOUND] No m:oMath or m:oMathPara nodes in document")

# Check for paragraphs with 'equation' style
print("\n--- Paragraphs with 'equation' style ---")
eq_style_paras = []
for i, p in enumerate(doc.paragraphs):
    if p.style and 'equation' in p.style.name.lower():
        eq_style_paras.append(i)
        text = p.text[:80]
        print(f"P{i:03d}: style={p.style.name}, text={text}")

if not eq_style_paras:
    print("[NONE FOUND] No paragraphs with 'equation' style")

# Check recent paragraphs around P068, P072, P077 (where equations were supposedly inserted)
print("\n--- Checking injection target areas ---")
for target_idx in [68, 69, 70, 71, 72, 73, 76, 77, 78]:
    if target_idx < len(doc.paragraphs):
        p = doc.paragraphs[target_idx]
        text = p.text[:60] if p.text else "[empty]"
        style = p.style.name if p.style else "[no style]"
        
        # Check XML structure
        has_omath = bool(p._p.findall('.//' + qn('m:oMath')))
        has_runs = bool(list(p._p.iter(W + 'r')))
        
        print(f"P{target_idx:03d}: style={style:15s} oMath={has_omath} runs={has_runs:5s} | {text}")

print("\n=== Summary ===")
print(f"Math nodes found: {len(math_nodes)}")
print(f"Equation-style paragraphs found: {len(eq_style_paras)}")
print(f"Expected: 9 equations (1)-(9)")
