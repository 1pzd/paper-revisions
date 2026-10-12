# -*- coding: utf-8 -*-
"""Fix duplicated equation-2 label in P37 (w:t node ')      (2)' -> ')'), restore pristine layout."""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

fixed = 0
for p in doc.paragraphs:
    if p.style.name != 'equation':
        continue
    for t in p._p.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
        if t.text and ')      (2)' in t.text:
            t.text = t.text.replace(')      (2)', ')')
            fixed += 1
            print('[fix] duplicate label removed in an equation paragraph')
if fixed == 0:
    print('[fix] nothing to do (already clean)')

doc.save(DOC)

# ---------- verify ----------
doc2 = Document(DOC)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'
M = '{http://schemas.openxmlformats.org/officeDocument/2006/math}t'
labels = []
for p in doc2.paragraphs:
    if p.style.name != 'equation':
        continue
    txt = ''.join((n.text or '') for n in p._p.iter(W)) + ''.join((n.text or '') for n in p._p.iter(M))
    found = re.findall(r'\((\d)\)', txt)
    print(f'[eq] nodes text tail ...{txt[-60:]!r} -> labels={found}')
    labels.extend(found)

print('[verify] labels:', labels)
ok_order = labels == [str(i) for i in range(1, 10)]
print(f'[verify] single labels 1..9 in order: {"PASS" if ok_order else "FAIL"}')

full = '\n'.join(p.text for p in doc2.paragraphs)
shapes = [(len(t.rows), len(t.columns)) for t in doc2.tables]
print('[verify] table shapes:', shapes)
ok_shapes = shapes == [(6, 2), (10, 5), (8, 7), (9, 5), (7, 7)]
ok_p5 = ('Limited and Imbalanced Data' in full and 'Not Label Rewriting' in full
         and 'Impact of Prompt Semantics (Row 7):' in full)
print(f'[verify] shapes: {"PASS" if ok_shapes else "FAIL"}; phase5 markers: {"PASS" if ok_p5 else "FAIL"}')
print(f'[verify] images: {len(doc2.inline_shapes)}')

if not (ok_order and ok_shapes and ok_p5):
    raise SystemExit('VERIFY FAIL')
print('[verify] PASS')
