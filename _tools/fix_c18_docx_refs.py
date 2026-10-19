# -*- coding: utf-8 -*-
"""
fix_c18_docx_refs.py — C18: delete 7 contaminated docx ref entries + renumber in-text [n] citations.

Authoritative user order (2026-10-17): "C18，删除污染条目并全文重编号"
Scope: CLIP.docx ONLY (txt ref list is already clean).

Deletes (body paras, ref-list region): Saunders drowning / Courvoisier / Tao EnFuseNet /
Jha Blue Books / Kim handpiece / PromptMix / composed image retrieval.

Old->new ref numbering (deletions: 2,7,8,9,16,19,20):
  1->1, 3->2, 4->3, 5->4, 6->5, 10->6, 11->7, 12->8, 13->9, 14->10, 15->11,
  17->12, 18->13, 21->14, 22->15, 23->16, 24->17, 25->18, 26->19

Special prose fixes (beyond pure renumbering, documented in CLAUDE.md):
  1. A17 'limited [19,20].' -> 'limited.'   (both refs deleted; general background claim kept uncited)
  2. A9  [7-10] collapses to [6] (Chen UPT); MOVE citation from paragraph tail to the
     long-tailed sentence it actually supports ('rare, severe cases[6].')
  3. A24 'LOCS III standard [13]' is a PRE-EXISTING MISCITATION (13=Dosovitskiy ViT;
     LOCS III = Chylack = old 3). After generic renumber ([13]->[9]) it is corrected to [2].

Backups: _backup/CLIP_pre_C18.docx (skip if exists).
"""
import os, re, shutil, sys
from docx import Document

DOCX = r'D:\修改论文\CLIP.docx'
BAK = r'D:\修改论文\_backup\CLIP_pre_C18.docx'
OUT = r'D:\修改论文\_tools\_dump_c18_result.txt'

log = []
def L(s):
    log.append(s)

# ---------- backup ----------
if not os.path.exists(BAK):
    shutil.copy2(DOCX, BAK)
    L('backup created: %s (%d bytes)' % (BAK, os.path.getsize(BAK)))
else:
    L('backup exists, skip: %s' % BAK)

doc = Document(DOCX)

# ---------- helpers ----------
def para_text(p):
    return ''.join(r.text or '' for r in p.runs)

def set_para_text(p, text):
    """Flatten para to its first run, preserving that run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ''

def replace_span_in_para(p, old, new):
    """Cross-run replace preserving run formatting. Index-based (p.runs returns
    fresh wrapper objects each access, so `is` identity checks must NOT be used)."""
    runs = p.runs
    full = ''.join(r.text or '' for r in runs)
    if old not in full:
        return False
    if full.count(old) > 1:
        raise RuntimeError('AMBIGUOUS span in para: %r' % old[:60])
    idx = full.index(old)
    end = idx + len(old)
    pos = 0
    si = ei = None
    so = eo = 0
    for j, r in enumerate(runs):
        t = r.text or ''
        if si is None and pos + len(t) > idx:
            si, so = j, idx - pos
        if pos < end <= pos + len(t):
            ei, eo = j, end - pos
            break
        pos += len(t)
    if si is None or ei is None:
        raise RuntimeError('span not coverable: %r' % old[:60])
    before = (runs[si].text or '')[:so]
    after = (runs[ei].text or '')[eo:]
    runs[si].text = before + new + after
    for j in range(si + 1, ei + 1):
        runs[j].text = ''
    return True

CITE_RE = re.compile(r'\[\d+(?:\s*[,;\-]\s*\d+)*\]')

def expand(token):
    'expand [4-6,16] -> [4,5,6,16]'
    inner = token[1:-1]
    nums = []
    for part in inner.split(','):
        part = part.strip()
        if '-' in part:
            a, b = part.split('-')
            nums.extend(range(int(a), int(b) + 1))
        else:
            nums.append(int(part))
    return nums

def compact(nums):
    'compact [4,5,6,16] -> [4-6,16]  (>=3 consecutive -> range, 2 -> comma pair)'
    if not nums:
        return ''
    nums = sorted(set(nums))
    parts = []
    i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        run = nums[i:j + 1]
        if len(run) >= 3:
            parts.append('%d-%d' % (run[0], run[-1]))
        else:
            parts.extend(str(x) for x in run)
        i = j + 1
    return '[' + ','.join(parts) + ']'

OLD2NEW = {1: 1, 3: 2, 4: 3, 5: 4, 6: 5, 10: 6, 11: 7, 12: 8, 13: 9, 14: 10,
           15: 11, 17: 12, 18: 13, 21: 14, 22: 15, 23: 16, 24: 17, 25: 18, 26: 19}
DELETED = {2, 7, 8, 9, 16, 19, 20}

# ---------- 1. delete 7 contaminated ref paragraphs ----------
DEL_KEYWORDS = [
    'Drowning Prevention',            # Saunders
    'Courvoisier',                    # abstract classification
    'EnFuseNet',                      # Tao skin lesion
    "Blue Books",                     # Jha tumor
    'active sentry and ozil handpiece',  # Kim
    'PromptMix',                      # Chen Y
    'composed image retrieval',       # Chen Z
]
deleted = []
for p in list(doc.paragraphs):
    t = para_text(p)
    for k in DEL_KEYWORDS:
        if k in t:
            deleted.append((k, t[:80]))
            p._element.getparent().remove(p._element)
            break
L('deleted ref paragraphs: %d' % len(deleted))
for k, t in deleted:
    L('  DEL [%s]: %s' % (k, t))
assert len(deleted) == 7, 'expected 7 deletions, got %d' % len(deleted)

# ---------- 2. special pre-fix: A17 drop [19,20] ----------
n = 0
for p in doc.paragraphs:
    if replace_span_in_para(p, 'limited [19,20].', 'limited.'):
        n += 1
L('pre-fix A17 drop [19,20]: %d' % n)
assert n == 1

# ---------- 3. generic renumber of all [n] tokens ----------
changed_tokens = 0
empty_dropped = 0
for p in doc.paragraphs:
    t = para_text(p)
    tokens = CITE_RE.findall(t)
    if not tokens:
        continue
    for tok in tokens:
        old_nums = expand(tok)
        new_nums = [OLD2NEW[x] for x in old_nums if x in OLD2NEW]
        dropped = [x for x in old_nums if x in DELETED]
        new_tok = compact(new_nums)
        if new_tok == tok and not dropped:
            continue
        if new_tok == '':
            # empty bracket: remove token + one preceding space if present
            if replace_span_in_para(p, ' ' + tok, ''):
                empty_dropped += 1
                changed_tokens += 1
                continue
            if replace_span_in_para(p, tok, ''):
                empty_dropped += 1
                changed_tokens += 1
                continue
            raise RuntimeError('cannot drop empty token %s' % tok)
        if replace_span_in_para(p, tok, new_tok):
            changed_tokens += 1
        else:
            raise RuntimeError('token replace failed %s -> %s' % (tok, new_tok))
L('generic renumber: %d tokens changed, %d empty dropped' % (changed_tokens, empty_dropped))

# ---------- 4. post-fix A9: move [6] to long-tail sentence ----------
n = 0
for p in doc.paragraphs:
    if 'interpretability[6].' in para_text(p):
        assert replace_span_in_para(p, 'rare, severe cases. Third', 'rare, severe cases[6]. Third')
        assert replace_span_in_para(p, 'interpretability[6].', 'interpretability.')
        n += 1
L('post-fix A9 move [6] to long-tail sentence: %d' % n)
assert n == 1

# ---------- 5. post-fix A24: LOCS III miscitation [9] -> [2] ----------
n = 0
for p in doc.paragraphs:
    if replace_span_in_para(p, 'LOCS III standard [9]', 'LOCS III standard [2]'):
        n += 1
L('post-fix A24 LOCS III miscitation [9]->[2]: %d' % n)
assert n == 1

# ---------- save ----------
doc.save(DOCX)
L('[docx] saved')

# ---------- VERIFY ----------
doc2 = Document(DOCX)
paras = doc2.paragraphs
alltxt = '\n'.join(para_text(p) for p in paras)
fails = []

# V1: exactly 19 ref entries after References heading
ref_texts = []
seen = False
for p in paras:
    t = para_text(p).strip()
    if t == 'References':
        seen = True
        continue
    if seen and t:
        ref_texts.append(t)
if len(ref_texts) != 19:
    fails.append('V1 ref count = %d (want 19)' % len(ref_texts))
else:
    L('V1 PASS: 19 ref entries')
if not ref_texts or 'World report on vision' not in ref_texts[0]:
    fails.append('V1b first ref is not WHO')
if not ref_texts or 'Mixup' not in ref_texts[-1]:
    fails.append('V1c last ref is not Mixup')

# V2: no contamination keywords anywhere
for k in DEL_KEYWORDS + ['Saunders J', 'Jha T', 'Tao Y']:
    if k in alltxt:
        fails.append('V2 contamination keyword present: %s' % k)
if not fails:
    L('V2 PASS: zero contamination keywords')

# V3: every cited number in 1..19; expected token map holds
all_toks = []
for p in paras:
    all_toks.extend(CITE_RE.findall(para_text(p)))
bad = [t for t in all_toks if any(x not in range(1, 20) for x in expand(t))]
if bad:
    fails.append('V3 out-of-range tokens: %s' % bad)
else:
    L('V3 PASS: all %d citation tokens within [1-19]: %s' % (len(all_toks), all_toks))

expect = ['[1]',                                     # A7
          '[2]', '[3-5]',                            # A8
          '[6]',                                     # A9 (moved to long-tail sentence)
          '[10,11]',                                 # A12
          '[7,8]', '[9]', '[4]', '[5]', '[10,11]', '[3-5]',   # A15
          '[12]', '[13]', '[14]', '[15]', '[16]', '[4,11]',   # A17 ([19,20] dropped)
          '[2]',                                     # A24 (miscite corrected to Chylack)
          '[10]',                                    # A25
          '[2]',                                     # A51
          '[17]',                                    # A78
          '[10,11]',                                 # A81
          '[18]',                                    # A91
          '[4,5,7-9,11,12]',                         # A94
          '[10]',                                    # A96
          '[12]', '[11]', '[4]', '[5]',              # A97
          '[19]']                                    # A323 conclusion
# order-sensitive full sequence check
if all_toks != expect:
    fails.append('V4 token sequence mismatch:\n got %s\nwant %s' % (all_toks, expect))
else:
    L('V4 PASS: full citation token sequence matches the renumber map')

# V5: CDRFF intact, no CDRF-Net
v5ok = True
if 'CDRFF-Net' not in alltxt:
    fails.append('V5 CDRFF-Net missing'); v5ok = False
if re.search(r'CDRF-Net', alltxt):
    fails.append('V5b bare CDRF-Net remains'); v5ok = False
if v5ok:
    L('V5 PASS: CDRFF-Net intact')

# V6: C4 four-limitations intact
if 'four concrete limitations' not in alltxt or 'BiomedCLIP' not in alltxt:
    fails.append('V6 C4 four-limitations paragraph damaged')
else:
    L('V6 PASS: C4 four-limitations paragraph intact')

# V7: conclusion cites [19]; A17 reads 'limited.' without [19,20]
if '[19]' not in para_text(paras[[i for i, p in enumerate(paras) if 'In conclusion' in para_text(p)][0]]):
    fails.append('V7 conclusion missing [19]')
if '[19,20]' in alltxt or '[20]' in alltxt:
    fails.append('V7b old [19,20] token remains')
else:
    L('V7 PASS: conclusion [19] present, no [19,20]/[20] leftovers')

# V8: no '[]' artifacts, no 'limited .' spacing damage
if '[]' in alltxt:
    fails.append('V8 empty [] artifact')
if 'limited .' in alltxt:
    fails.append('V8b spacing damage around removed citation')
else:
    L('V8 PASS: no bracket/spacing artifacts')

L('VERIFY: %d/8 PASS' % (8 - len(fails)) if not fails else 'VERIFY FAILURES:\n' + '\n'.join(fails))

open(OUT, 'w', encoding='utf-8').write('\n'.join(log))
print('C18 done. %s' % ('ALL PASS' if not fails else 'FAILURES'))
sys.exit(1 if fails else 0)
