import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

print('table shapes:', [(len(t.rows), len(t.columns)) for t in doc.tables])

full = '\n'.join(p.text for p in doc.paragraphs)
missing = [i for i in range(1, 10) if f'({i})' not in full]
print('missing (i) labels in p.text:', missing)

# where are equation paragraphs and what do their texts look like?
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'equation':
        runs = [r.text for r in p.runs]
        print(f'P{i} [equation] p.text={p.text!r} runs={runs!r}')

# raw XML text of one equation paragraph (first)
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'equation':
        xml = p._p.xml
        ts = re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml)
        print(f'P{i} all w:t nodes: {ts!r}')
        break

# search labels anywhere including tables
print('labels in tables:', [(i, [c.text for r in t.rows for c in r.cells if re.search(r'\(\d\)', c.text)]) for i, t in enumerate(doc.tables)])
