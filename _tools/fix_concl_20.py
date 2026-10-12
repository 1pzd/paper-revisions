# -*- coding: utf-8 -*-
"""Re-attach [20] Mixup citation to a legitimate future-work clause (zero orphan refs)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

concl = None
for p in doc.paragraphs:
    if p.text.startswith('In conclusion, this study presents CLIP-CAT'):
        concl = p
        break
if concl is None:
    raise RuntimeError('conclusion not found')

if 'mixup-style augmentation' in concl.text:
    print('[guard] already applied')
else:
    r0, r1, r2 = concl.runs[0], concl.runs[1], concl.runs[2]
    if r1.text not in ('', '[20]'):
        raise RuntimeError(f'unexpected run1 {r1.text!r}')
    tail = 'and refining prompt engineering to enhance robustness against imaging artifacts and complex ocular conditions in real-world scenarios'
    if not r0.text.endswith(tail):
        raise RuntimeError('run0 tail unexpected')
    r0.text = (r0.text[:-len(tail)]
               + 'refining prompt engineering, and exploring mixup-style augmentation')
    r1.text = '[20]'
    r2.text = ' to enhance robustness against imaging artifacts and complex ocular conditions in real-world scenarios.'
    doc.save(DOC)
    print('[saved] conclusion [20] re-attached')

doc2 = Document(DOC)
for p in doc2.paragraphs:
    if p.text.startswith('In conclusion, this study presents CLIP-CAT'):
        print('[verify]', p.text[-260:])
        assert 'mixup-style augmentation[20]' in p.text, 'citation not adjacent to clause'
        assert 'scenarios[20]' not in p.text
        break
print('[verify] PASS')
