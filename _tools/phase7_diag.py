# -*- coding: utf-8 -*-
"""Investigate: Table 1 full dump + proper eq-label check (reuse final_check logic)."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

print('=== Table 1 (tables[0]) full dump ===')
t1 = doc.tables[0]
print(f'shape: {len(t1.rows)}x{len(t1.columns)}')
for i, row in enumerate(t1.rows):
    for j, cell in enumerate(row.cells):
        txt = cell.text
        print(f'  r{i}c{j}: {txt[:120]!r}')

print()
print('=== eq labels: paragraphs containing m:oMath ===')
labels = []
for idx, p in enumerate(doc.paragraphs):
    has_math = len(list(p._p.iter(qn('m:oMath')))) > 0
    if not has_math:
        continue
    # collect all text nodes w:t + m:t
    texts = [n.text or '' for n in p._p.iter(qn('m:t'))]
    texts += [n.text or '' for n in p._p.iter(qn('w:t'))]
    digits = [t.strip() for t in texts if t.strip() in list('123456789')]
    print(f'  P{idx:03d} style={p.style.name!r} math_nodes={len(list(p._p.iter(qn("m:oMath"))))} digit_nodes={digits}')
    if digits:
        labels.append(digits[-1])  # label is the trailing digit node
print(f'[labels] {labels}')
print(f'[check] labels == 1..9: {labels == list("123456789")}')
