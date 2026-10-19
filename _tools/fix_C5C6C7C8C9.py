# -*- coding: utf-8 -*-
"""C5-C9 fixes for CLIP.docx + CLIP-revised.txt (2026-10-15).
C5: compress Related Work 2.2 CLIP catalog + add ophthalmology clinical refs.
C6: reposition novelty as combination contribution (L_rank = expected |distance|).
C7: define CLIP Base configuration.
C8: clarify Imbalance Strategy switch semantics.
C9: unify Bidirectional Rank-Aware Loss terminology.
"""
import shutil, sys, io, os
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "CLIP.docx")
TXT = os.path.join(ROOT, "CLIP-revised.txt")
BACK = os.path.join(ROOT, "_backup")

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

# ---------------- backups ----------------
shutil.copy2(DOCX, os.path.join(BACK, "CLIP_pre_C5C9.docx"))
shutil.copy2(TXT, os.path.join(BACK, "CLIP-revised_pre_C5C9.txt"))
print("[backup] done")

# ================= TXT =================
with open(TXT, encoding="utf-8") as f:
    lines = f.readlines()
text = "".join(lines)

def rep_t(old, new, tag):
    global text
    if old not in text:
        raise SystemExit("TXT MISS [%s]: %r" % (tag, old[:80]))
    if text.count(old) != 1:
        raise SystemExit("TXT AMBIG [%s]: count=%d" % (tag, text.count(old)))
    text = text.replace(old, new)
    print("[txt] ok", tag)

# C9/C6 abstract
rep_t("we design a Bidirectional Rank-Aware Loss that encourages unimodal",
      "we incorporate a Bidirectional Rank-Aware Loss that encourages unimodal",
      "abs-design")

# C6+C9 contribution 2
rep_t("Second, we propose a Bidirectional Rank-Aware Constraint that redefines the classification loss. By enforcing unimodal probability distributions and penalizing distance-based errors, we ensure predictions respect the strict logical progression of disease severity.",
      "Second, we incorporate a Bidirectional Rank-Aware Loss into the classification objective. The loss penalizes the expected absolute distance between the predicted distribution and the ground-truth grade, placing it in the established family of expected-distance ordinal objectives (7,8); its bidirectional character reflects the symmetry of the absolute distance, which charges under- and overestimation alike. The contribution is therefore combinatorial: this rank objective is coupled with language-aligned semantic anchors and a multi-granularity encoder so that predictions respect the strict logical progression of disease severity.",
      "contrib2")

# C5 related work 2.2 rewrite (full line)
old_l41_prefix = "Vision-language models such as CLIP (12) align"
l41_idx = None
for i, ln in enumerate(lines):
    if ln.startswith(old_l41_prefix):
        l41_idx = i
        break
if l41_idx is None:
    raise SystemExit("TXT MISS l41 anchor")
new_l41 = (
    "Vision-language models such as CLIP (12) align images and text in a shared embedding space. "
    "Subsequent work scales contrastive pretraining and data pipelines (13-23), extends objectives toward multimodal generation and matching (24-26), and develops prompt-based and adapter-based adaptation for downstream transfer (27-40). "
    "In medicine, these models have mainly been applied in radiology and pathology, while ophthalmic adoption remains limited. In cataract assessment, deep grading systems built on clinical scales have been developed for slit-lamp and retro-illumination photographs (41), nuclear cataract diagnosis from slit-lamp video (42), and cortical cataract grading with clinical sub-scores (43), but these pipelines remain CNN- or Transformer-based classifiers without language-aligned semantics. "
    "Standard contrastive alignment, moreover, does not preserve the ordinal topology of cataract severity. Unlike prior pipelines that attach ordinal losses to convolutional neural network features (4,7,8), this study couples vision-language alignment with a Bidirectional Rank-Aware Loss and grade-specific medical prompts.\n"
)
text = text.replace(lines[l41_idx], new_l41, 1)
print("[txt] ok l41-related-work")

# C9 framework component name
rep_t("Bidirectional Rank-Aware Optimization: To address the ordinal nature of cataracts, we introduce a specialized loss function that penalizes predictions based on their distance from the true grade, enforcing a unimodal probability distribution.",
      "Bidirectional Rank-Aware Loss: To address the ordinal nature of cataracts, we integrate a distance-weighted rank loss that penalizes predictions based on their expected ordinal distance from the true grade, encouraging a unimodal probability distribution.",
      "comp3-name")

# C6+C8+C9 sec 3.4
rep_t("The complete training objective combines a focal term and a bidirectional rank-aware term.",
      "The complete training objective combines a focal term and the Bidirectional Rank-Aware Loss (hereafter the rank loss).",
      "34-obj")
rep_t("The bidirectional rank-aware term penalizes the expected ordinal distance:",
      "The rank loss penalizes the expected ordinal distance:",
      "34-lrank")
rep_t("The first sum charges probability mass placed on grades below the true grade (underestimation), and the second charges mass placed above it (overestimation), both weighted linearly by grade distance.",
      "The first sum charges probability mass placed on grades below the true grade (underestimation), and the second charges mass placed above it (overestimation), both weighted linearly by grade distance. By construction, L_rank equals the expected absolute distance between the predicted distribution and the ground-truth grade, placing it in the established family of expected-distance ordinal objectives (7,8); its bidirectional character reflects the symmetry of the absolute distance, which charges under- and overestimation alike, rather than a new loss form.",
      "34-c6")
rep_t("The total loss is L_total = L_focal + λ·L_rank, where λ=1.0 balances the two terms.",
      "The total loss is L_total = L_focal + λ·L_rank, where λ=1.0 balances the two terms. In the ablation study (Table 2), the Imbalance Strategy switch jointly enables class-weighted sampling and the focal term; rows without it train with uniform sampling and the standard cross-entropy term in place of L_focal.",
      "34-c8")

# C7 ablation baseline sentence
rep_t("The naive application of CLIP yields an accuracy of 80.15%.",
      "The CLIP Base configuration (Row 1)—the global CLIP ViT-L/14 branch with grade-specific medical prompts, fine-tuned as described in Section 3.6 with the local branch, rank loss, and imbalance handling all disabled—yields an accuracy of 80.15%, distinct from the zero-shot (64.20%) and linear-probe (79.50%) transfer settings in Table 1.",
      "c7-baseline")
rep_t("The integration of Rank Loss provides",
      "The integration of the rank loss provides",
      "c9-rankloss-integration")

# C9 discussion
rep_t("The bidirectional rank-aware loss enforces",
      "The Bidirectional Rank-Aware Loss enforces",
      "c9-discussion")

# C7/C8 table caption
rep_t("Table 2. Ablation study demonstrating the incremental improvements of each module.",
      "Table 2. Ablation study demonstrating the incremental improvements of each module. Rank Loss denotes the Bidirectional Rank-Aware Loss; Imbalance Strategy denotes the joint use of class-weighted sampling and the focal loss term (rows marked '-' train with uniform sampling and standard cross-entropy); CLIP Base denotes the fine-tuned global CLIP ViT-L/14 branch with medical prompts (Section 3.6), distinct from the zero-shot and linear-probe rows of Table 1.",
      "c7c8-caption")

# C9 intro + fig1 + fig2 lowercase variants
rep_t("with a bidirectional rank-aware loss, our approach outperforms",
      "with the Bidirectional Rank-Aware Loss, our approach outperforms",
      "c9-intro")
rep_t("trained with the bidirectional rank-aware loss",
      "trained with the Bidirectional Rank-Aware Loss",
      "c9-fig1")
rep_t("of the rank-aware loss",
      "of the rank loss",
      "c9-fig2")

# reference renumber: old (41)(42)(43) -> (44)(45)(46) at known sites
rep_t("dominate the gradient (41).",
      "dominate the gradient (44).",
      "ref41-shift")
rep_t("AdamW optimizer (42)",
      "AdamW optimizer (45)",
      "ref42-shift")
rep_t("mixup-style augmentation (43)",
      "mixup-style augmentation (46)",
      "ref43-shift")

# reference list: renumber old 41-43 headers, insert new 41-43
old_refs_block = None
for i, ln in enumerate(lines):
    if ln.startswith("41. Lin T Y"):
        old_refs_block = i
        break
if old_refs_block is None:
    # find any '41.' reference line
    for i, ln in enumerate(lines):
        if ln.startswith("41. "):
            old_refs_block = i
            break
if old_refs_block is None:
    raise SystemExit("TXT MISS ref41 list line")

new_ref_lines = [
    "41. Son KY, Ko J, Kim E, Lee SY, Kim MJ, Han J, et al. Deep learning-based cataract detection and grading from slit-lamp and retro-illumination photographs. Ophthalmol Sci (2022). doi: 10.1016/j.xops.2022.100147\n",
    "42. Shimizu E, Tanji M, Nakayama S, Ishikawa T, Agata N. AI-based diagnosis of nuclear cataract from slit-lamp videos. Sci Rep (2023). doi: 10.1038/s41598-023-49563-7\n",
    "43. Wang J, Xu Z, Zheng W, Ying H, Chen T, Liu Z, et al. A transformer-based knowledge distillation network for cortical cataract grading. IEEE Trans Med Imaging (2023). doi: 10.1109/TMI.2023.3327274\n",
]

# re-derive current text lines after earlier edits
lines2 = text.splitlines(keepends=True)
ref41_idx = None
for i, ln in enumerate(lines2):
    if ln.startswith("41. "):
        ref41_idx = i
        break
if ref41_idx is None:
    raise SystemExit("TXT MISS ref41 line after edits")
def renum(prefix_old, prefix_new):
    for i in range(ref41_idx, min(ref41_idx + 12, len(lines2))):
        if lines2[i].startswith(prefix_old):
            lines2[i] = prefix_new + lines2[i][len(prefix_old):]
            return
    raise SystemExit("TXT MISS renumber %s" % prefix_old)
renum("43. ", "46. ")
renum("42. ", "45. ")
renum("41. ", "44. ")
lines2[ref41_idx:ref41_idx] = new_ref_lines
text = "".join(lines2)
print("[txt] ok refs-insert-41-43")

with open(TXT, "w", encoding="utf-8", newline="") as f:
    f.write(text)
print("[txt] saved")

# ================= DOCX =================
doc = Document(DOCX)
paras = doc.paragraphs

def find_para(substr, start=0):
    for i in range(start, len(doc.paragraphs)):
        if substr in doc.paragraphs[i].text:
            return i
    raise SystemExit("DOCX MISS: %r" % substr[:60])

def has_math(p):
    return "<m:oMath" in p._p.xml

def rep_p(p, old, new, tag):
    # run-local first
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            print("[docx] ok(run)", tag)
            return
    if old in p.text:
        if has_math(p):
            raise SystemExit("DOCX MATH-WRITE BLOCKED [%s]" % tag)
        # rebuild preserving first run pr
        runs = p.runs
        if not runs:
            raise SystemExit("DOCX NO-RUN [%s]" % tag)
        full = p.text.replace(old, new)
        runs[0].text = full
        for r in runs[1:]:
            r.text = ""
        print("[docx] ok(rebuild)", tag)
        return
    raise SystemExit("DOCX MISS [%s]: %r" % (tag, old[:60]))

# C6/C9 abstract
i = find_para("Bidirectional Rank-Aware Loss")
try:
    rep_p(doc.paragraphs[i], "we design a Bidirectional Rank-Aware Loss",
          "we incorporate a Bidirectional Rank-Aware Loss", "abs-design")
except SystemExit as e:
    if "MISS" in str(e):
        print("[docx] abstract already uses different phrasing; skip")
    else:
        raise

# C6+C9 contribution (P012 area)
i = find_para("We propose a Bidirectional Rank-Aware Constraint")
rep_p(doc.paragraphs[i],
      "We propose a Bidirectional Rank-Aware Constraint that redefines the classification loss. By enforcing unimodal probability distributions and penalizing distance-based errors, we ensure predictions respect the strict logical progression of disease severity.",
      "We incorporate a Bidirectional Rank-Aware Loss into the classification objective. The loss penalizes the expected absolute distance between the predicted distribution and the ground-truth grade, placing it in the established family of expected-distance ordinal objectives [14,15]; its bidirectional character reflects the symmetry of the absolute distance, which charges under- and overestimation alike. The contribution is therefore combinatorial: this rank objective is coupled with language-aligned semantic anchors and a multi-granularity encoder so that predictions respect the strict logical progression of disease severity.",
      "contrib2")

# C5 related work (P017)
i = find_para("yet ophthalmic adoption remains limited")
rep_p(doc.paragraphs[i],
      "yet ophthalmic adoption remains limited [19,20]. Standard contrastive alignment",
      "yet ophthalmic adoption remains limited [19,20]. In cataract assessment, deep grading systems built on clinical scales have been developed for slit-lamp and retro-illumination photographs [21], nuclear cataract diagnosis from slit-lamp video [22], and cortical cataract grading with clinical sub-scores [23], but these pipelines remain CNN- or Transformer-based classifiers without language-aligned semantics. Standard contrastive alignment",
      "c5-clinical")
i = find_para("couples vision-language alignment with a bidirectional rank-aware objective")
rep_p(doc.paragraphs[i],
      "with a bidirectional rank-aware objective and grade-specific medical prompts",
      "with a Bidirectional Rank-Aware Loss and grade-specific medical prompts",
      "c5-tail")

# C9 framework component (P023)
i = find_para("Bidirectional Rank-Aware Optimization")
rep_p(doc.paragraphs[i],
      "Bidirectional Rank-Aware Optimization: To address the ordinal nature of cataracts, we introduce a specialized loss function that penalizes predictions based on their distance from the true grade, enforcing a unimodal probability distribution.",
      "Bidirectional Rank-Aware Loss: To address the ordinal nature of cataracts, we integrate a distance-weighted rank loss that penalizes predictions based on their expected ordinal distance from the true grade, encouraging a unimodal probability distribution.",
      "comp3-name")

# C9 sec 3.4 term def (P067)
i = find_para("The bidirectional rank-aware term penalizes")
rep_p(doc.paragraphs[i],
      "The bidirectional rank-aware term penalizes",
      "The Bidirectional Rank-Aware Loss (hereafter the rank loss) penalizes",
      "34-name")

# C6 append after rank penalty paragraph (P069)
i = find_para("Minimizing the rank term therefore concentrates")
p = doc.paragraphs[i]
if has_math(p):
    p.add_run(" By construction, the rank loss equals the expected absolute distance between the predicted distribution and the ground-truth grade, a construction in the same family as expected-distance ordinal objectives [14,15]; the term “bidirectional” refers to the symmetry of the absolute distance rather than to a new loss family.")
else:
    p.add_run(" By construction, the rank loss equals the expected absolute distance between the predicted distribution and the ground-truth grade, a construction in the same family as expected-distance ordinal objectives [14,15]; the term “bidirectional” refers to the symmetry of the absolute distance rather than to a new loss family.")
print("[docx] ok append c6")

# C8 append to Focal Loss paragraph (P071)
i = find_para("This scheme utilizes Focal Loss")
p = doc.paragraphs[i]
p.add_run(" In the ablation study, the Imbalance Strategy switch jointly enables class-weighted sampling and the focal term; rows without it train with uniform sampling and the standard cross-entropy term in place of the focal term.")
print("[docx] ok append c8")

# C7 baseline paragraph (P089)
i = find_para("Baseline (Row 1)")
rep_p(doc.paragraphs[i],
      "Baseline (Row 1): The naive application of CLIP yields an accuracy of 80.15%.",
      "Baseline (Row 1): The CLIP Base configuration—the global CLIP ViT-L/14 branch with grade-specific medical prompts, fine-tuned as described in Implementation Details with the local branch, rank loss, and imbalance handling all disabled—yields an accuracy of 80.15%, distinct from the zero-shot (64.20%) and linear-probe (79.50%) transfer settings in Table 2.",
      "c7-baseline")

# C9 rank loss integration (P091)
i = find_para("Impact of Rank Loss")
rep_p(doc.paragraphs[i],
      "Impact of Rank Loss (Row 3): The integration of Rank Loss provides",
      "Impact of the Rank Loss (Row 3): The integration of the rank loss provides",
      "c9-impact")

# C9 fig captions
i = find_para("trained with the bidirectional rank-aware loss")
rep_p(doc.paragraphs[i],
      "trained with the bidirectional rank-aware loss",
      "trained with the Bidirectional Rank-Aware Loss",
      "c9-fig1")
i = find_para("of the rank-aware loss")
rep_p(doc.paragraphs[i],
      "of the rank-aware loss",
      "of the rank loss",
      "c9-fig2")

# C7/C8 table 3 caption
i = find_para("Ablation study demonstrating the incremental improvements")
p = doc.paragraphs[i]
p.add_run(" Rank Loss denotes the Bidirectional Rank-Aware Loss; Imbalance Strategy denotes the joint use of class-weighted sampling and the focal loss term (rows without it train with uniform sampling and standard cross-entropy); CLIP Base denotes the fine-tuned global CLIP ViT-L/14 branch with medical prompts, distinct from the zero-shot and linear-probe rows of Table 2.")
print("[docx] ok append caption")

# renumber [21]->[24] etc (run-level; P079 has math)
i = find_para("gradient[21]")
rep_p(doc.paragraphs[i], "[21]", "[24]", "ref21")
i = find_para(")[22]")
rep_p(doc.paragraphs[i], "[22]", "[25]", "ref22")
i = find_para("augmentation[23]")
rep_p(doc.paragraphs[i], "[23]", "[26]", "ref23")

# insert 3 reference items before Lin focal loss entry
i = find_para("Focal loss for dense object detection")
anchor = doc.paragraphs[i]
style = anchor.style
new_refs = [
    "Son K Y ,Ko J ,Kim E , et al. Deep Learning-Based Cataract Detection and Grading from Slit-Lamp and Retro-Illumination Photographs[J].Ophthalmology Science,2022. DOI:10.1016/J.XOPS.2022.100147",
    "Shimizu E ,Tanji M ,Nakayama S , et al. AI-based diagnosis of nuclear cataract from slit-lamp videos[J].Scientific Reports,2023. DOI:10.1038/S41598-023-49563-7",
    "Wang J ,Xu Z ,Zheng W , et al. A Transformer-Based Knowledge Distillation Network for Cortical Cataract Grading[J].IEEE Transactions on Medical Imaging,2023. DOI:10.1109/TMI.2023.3327274",
]
for t in new_refs:
    np_ = anchor.insert_paragraph_before(t, style=style)
print("[docx] ok insert-3-refs")

doc.save(DOCX)
print("[docx] saved")

# ================= VERIFY =================
doc2 = Document(DOCX)
n_para = len(doc2.paragraphs)
n_tbl = len(doc2.tables)
n_math = sum(p._p.xml.count("<m:oMath") for p in doc2.paragraphs)
n_ref = sum(1 for p in doc2.paragraphs if p.style.name == "referenceitem")
print("[verify] docx paras=%d tables=%d oMath=%d refitems=%d" % (n_para, n_tbl, n_math, n_ref))
alltext = "\n".join(p.text for p in doc2.paragraphs)
for k in ["Son K Y", "Shimizu", "Wang J", "expected absolute distance",
          "hereafter the rank loss", "[24]", "[25]", "[26]", "CLIP Base configuration",
          "Bidirectional Rank-Aware Constraint", "Bidirectional Rank-Aware Optimization"]:
    print("[verify] docx %-32s %s" % (k, "FOUND" if k in alltext else "ABSENT"))
t2 = open(TXT, encoding="utf-8").read()
for k in ["Son KY", "Shimizu E", "Wang J", "expected absolute distance",
          "hereafter the rank loss", "(44)", "(45)", "(46)", "CLIP Base configuration",
          "Bidirectional Rank-Aware Constraint", "Bidirectional Rank-Aware Optimization"]:
    print("[verify] txt  %-32s %s" % (k, "FOUND" if k in t2 else "ABSENT"))
print("DONE")
