# -*- coding: utf-8 -*-
"""Grammar fix in conclusion future-work clause."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)
for p in doc.paragraphs:
    if p.text.startswith('In conclusion, this study presents CLIP-CAT'):
        r0 = p.runs[0]
        if 'external validation refining prompt engineering,' in r0.text:
            r0.text = r0.text.replace('external validation refining prompt engineering,',
                                      'external validation, refining prompt engineering,')
            doc.save(DOC)
            print('[saved] grammar fixed')
        elif 'external validation, refining prompt engineering,' in r0.text:
            print('[guard] already fixed')
        else:
            raise RuntimeError('anchor not found')
        break

doc2 = Document(DOC)
for p in doc2.paragraphs:
    if p.text.startswith('In conclusion, this study presents CLIP-CAT'):
        tail = p.text[p.text.find('Future work'):]
        print('[verify]', tail)
        assert 'external validation, refining prompt engineering, and exploring mixup-style augmentation[20]' in p.text
        break
print('[verify] PASS')
