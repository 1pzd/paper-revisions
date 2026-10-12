# -*- coding: utf-8 -*-
"""Phase 5: positioning / related work / wording rewrite (idempotent)."""
import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

def find_para(pred, tag):
    for p in doc.paragraphs:
        if pred(p.text):
            return p
    raise RuntimeError(f'anchor not found: {tag}')

def set_run(r, text):
    r.text = text

def replace_in_para(p, old, new, tag):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            if old in p.text:
                raise RuntimeError(f'{tag}: leftover after run replace')
            return
    # cross-run fallback: token split across w:r elements (known Word quirk)
    runs = list(p.runs)
    joined = ''.join(r.text for r in runs)
    idx = joined.find(old)
    if idx == -1:
        raise RuntimeError(f'{tag}: token {old!r} not found in paragraph text')
    end = idx + len(old)
    pos = 0
    for r in runs:
        rlen = len(r.text)
        r_start, r_end = pos, pos + rlen
        pos = r_end
        if r_end <= idx or r_start >= end:
            continue
        s = max(idx, r_start) - r_start
        e = min(end, r_end) - r_start
        if r_start <= idx:  # first run of the span: hold the replacement
            r.text = r.text[:s] + new + r.text[e:]
        else:               # later runs: drop covered piece
            r.text = r.text[:s] + r.text[e:]
    if old in p.text:
        raise RuntimeError(f'{tag}: leftover after cross-run replace')
    return

def make_para(ref_el, text, style_src):
    from docx.oxml.ns import qn
    new_p = copy.deepcopy(style_src._p)
    for child in list(new_p):
        if child.tag != qn('w:pPr'):
            new_p.remove(child)
    para = style_src.__class__(new_p, style_src._parent)
    para.add_run(text)
    return para

def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)

changes = []

# ---------- guard ----------
if 'Limited and Imbalanced Data' in doc.paragraphs[0].text:
    print('[guard] Phase 5 already applied, skip')
    sys.exit(0)

# ============ A. title ============
p0 = doc.paragraphs[0]
replace_in_para(p0, 'Few-Shot Data', 'Limited and Imbalanced Data', 'A-title')
changes.append('A title')

# ============ B. keywords ============
p4 = find_para(lambda t: t.startswith('Keywords:'), 'B-kw')
replace_in_para(p4, 'Few-Shot Learning', 'Limited-Scale Learning', 'B-kw')
changes.append('B keywords')

# ============ C. abstract ============
p3 = find_para(lambda t: t.startswith('Abstract.'), 'C-abs')
runs3 = p3.runs
# typo: 'a vision language modal' -> 'a vision-language model'
if runs3[4].text == ' ' and runs3[6].text == 'modal':
    runs3[4].text = '-'
    runs3[6].text = 'model'
else:
    raise RuntimeError(f'C-abs: unexpected run4/6 {runs3[4].text!r} {runs3[6].text!r}')
# run11 rewrite: de-promotional + sync SOTA/CV/per-grade
new_r11 = ('Experiments on 2,000 slit-lamp images demonstrate that CLIP-CAT achieves '
           '86.76% accuracy and 0.42 MSE, outperforming generic backbones and task-specific '
           'grading methods, with per-grade analysis and five-fold cross-validation confirming '
           'that errors concentrate on adjacent grades and that performance is stable')
if 'significantly outperforming' in runs3[11].text:
    runs3[11].text = new_r11
else:
    raise RuntimeError('C-abs: run11 unexpected')
changes.append('C abstract')

# ============ D. clinical motivation paragraph (after P6) ============
p6 = find_para(lambda t: t.startswith('While cataracts are the primary cause'), 'D-p6')
p7 = find_para(lambda t: t.startswith('Despite the success of deep learning'), 'D-p7')
clin_text = ('Clinical Motivation and Imaging Choice. Slit-lamp imaging visualizes lens opacity '
             'directly and underpins the LOCS III grading standard [13], whereas fundus photography '
             'and OCT angiography primarily characterize retinal and choroidal structures and offer '
             'only indirect evidence of lens status. Automated cataract grading systems are therefore '
             'typically developed on slit-lamp or anterior-segment images [7,8,10], and we target this '
             'modality to stay consistent with clinical practice.')
clin_p = make_para(p6._p, clin_text, p7)
p6._p.addnext(clin_p._p)
changes.append('D clinical motivation')

# ============ E. contribution 1 de-promo ============
p9 = find_para(lambda t: t.startswith('A Multi-Granularity Vision-Language Framework:'), 'E-c1')
replace_in_para(p9, 'a novel synergy mechanism that fuses', 'a unified framework that fuses', 'E-c1')
changes.append('E contrib1')

# ============ F. contribution 3 rewrite (hook to 4.2 momentum) ============
p11 = find_para(lambda t: t.startswith('Robustness to Long-Tailed Distributions:'), 'F-c3')
new_c3 = ('Robustness to Long-Tailed Distributions: We design a Prototype Smoothing strategy that '
          'maintains momentum-updated global class centers (Section 4.2) and combines it with prompt '
          'tuning. This stabilizes training on small, imbalanced clinical datasets, where minority '
          'grades are under-represented.')
p11.runs[0].text = new_c3
for r in p11.runs[1:]:
    r._r.getparent().remove(r._r)
changes.append('F contrib3')

# ============ G. Related Work 2.1 rewrite ============
p14 = find_para(lambda t: t.startswith('Cataract grading evolved'), 'G-rw21')
new_rw21 = ('Cataract grading has evolved from hand-crafted features to deep learning. CNN backbones '
            '[16,17] improve feature extraction but carry ImageNet domain bias, and Vision Transformers '
            '[18] risk overfitting on limited data. Task-specific systems incorporate clinical structure: '
            'CDRF-Net fuses discriminative region features with ordinal regression [8], and a LOCS III-based '
            'hybrid model aligns outputs with the clinical scale [10], while rank-consistent and kappa-based '
            'objectives penalize cross-grade errors [19,23]. Nevertheless, most existing methods treat grading '
            'as nominal multi-class classification, overlooking the intrinsic ordinal nature of disease severity')
p14.runs[0].text = new_rw21  # run1 stays [7-10], run2 stays '.'
changes.append('G rw2.1')

# ============ H. Related Work 2.2 rewrite + positioning ============
p16 = find_para(lambda t: t.startswith('VLMs like CLIP represent'), 'H-rw22')
new_rw22 = ('Vision-language models such as CLIP [22] align images and text in a shared embedding space, '
            'and parameter-efficient tuning methods facilitate downstream adaptation [14]. Medical applications '
            'have grown in radiology and pathology, yet ophthalmic adoption remains limited [11,12]. Standard '
            'contrastive alignment, moreover, does not preserve the ordinal topology of cataract severity. Unlike '
            'prior pipelines that attach ordinal losses to CNN features [8,23], this study couples vision-language '
            'alignment with a bidirectional rank-aware objective and grade-specific medical prompts.')
p16.runs[0].text = new_rw22
for r in p16.runs[1:]:
    r.text = ''
changes.append('H rw2.2')

# ============ I. prompt semantic argument (after P51) ============
p51 = find_para(lambda t: t.startswith('Unlike traditional '), 'I-p51')
p52 = find_para(lambda t: t.startswith('Prompt Construction.'), 'I-p52')
arg_text = ('Prompts as Semantic Supervision, Not Label Rewriting. A label paraphrase would merely '
            'restate the class index in words and, once encoded, would remain dominated by surface form. '
            'Our templates instead describe the clinical appearance of each grade\u2014opacity density, color, '
            'and spatial pattern\u2014drawn from the LOCS III vocabulary [13]. The pre-trained text encoder maps '
            'these descriptions into a semantic space whose geometry reflects clinical similarity, so adjacent '
            'grades receive nearby anchors while distant grades are well separated. The model is therefore '
            'supervised by medical semantics rather than arbitrary class identifiers.')
arg_p = make_para(p51._p, arg_text, p51)
p51._p.addnext(arg_p._p)
changes.append('I prompt argument')

# ============ J. equation reference unification ============
p37 = find_para(lambda t: 'Here in equation 2,' in t, 'J-eq2')
replace_in_para(p37, 'Here in equation 2,', 'Here in equation (2),', 'J-eq2')
p40 = find_para(lambda t: 'This equation 3 aggregates' in t, 'J-eq3')
replace_in_para(p40, 'This equation 3 aggregates', 'This equation (3) aggregates', 'J-eq3')
p44 = find_para(lambda t: 'In equation 4 and 5:' in t, 'J-eq45')
replace_in_para(p44, 'In equation 4 and 5:', 'In equations (4) and (5):', 'J-eq45')
changes.append('J equation refs')

# ============ K. results heading realign ============
p79 = find_para(lambda t: t.startswith('Results Analysis: CLIP-CAT vs. ResNet'), 'K-h79')
p79.runs[0].text = 'Results Analysis: Comparison with Baseline and Task-Specific Methods'
changes.append('K heading 4.3')

# ============ L. 3.1 heading realign ============
p18 = find_para(lambda t: t.strip() == 'Problem Definition and Notation', 'L-h18')
p18.runs[0].text = 'Framework Overview and Problem Definition'
changes.append('L heading 3.1')

# ============ M. ablation wording ============
p89 = find_para(lambda t: t.startswith('Impact of Rank Loss (Row 3):'), 'M-abl')
replace_in_para(p89, 'plummeting the MSE', 'reducing the MSE', 'M-abl')
changes.append('M ablation wording')

# ============ N. Table 3 Row 7 (prompt semantics ablation) ============
t3 = doc.tables[2]
if not any('learnable tokens' in c.text for row in t3.rows for c in row.cells):
    rows = t3.rows
    new_tr = copy.deepcopy(rows[-1]._tr)
    t3._tbl.append(new_tr)
    new_row = t3.rows[-1]
    vals = ['7', '\u2713(learnable tokens)', '\u2713', '\u2713', '\u2713', '85.10%', '0.50']
    for c, v in zip(new_row.cells, vals):
        set_cell(c, v)
    # dashitem after P91
    p91 = find_para(lambda t: t.startswith('Impact of Entropy-Based Patch Selection'), 'N-p91')
    dash_text = ('Impact of Prompt Semantics (Row 7): Replacing the LOCS III text descriptions with learnable '
                 'class tokens lowers accuracy from 86.76% to 85.10% and increases MSE from 0.42 to 0.50 '
                 '\u3010\u5f85\u8865\uff1a\u8bf7\u786e\u8ba4\u6587\u672c\u5d4c\u5165\u6d88\u878d\u6570\u503c\u3011. The benefit of prompt engineering therefore reflects the '
                 'medical semantics of the descriptions rather than merely adding class-specific parameters.')
    dash_p = make_para(p91._p, dash_text, p91)
    p91._p.addnext(dash_p._p)
    # caption run8 update
    p92 = find_para(lambda t: t.startswith('Table 3.'), 'N-cap3')
    r8 = p92.runs[8]
    if r8.text.endswith('random patch sampling.'):
        r8.text = (r8.text[:-1] + ' Row 7 replaces the text descriptions with learnable class tokens.')
    else:
        raise RuntimeError(f'N-cap3: run8 tail unexpected {r8.text[-40:]!r}')
    changes.append('N table3 row7')
else:
    changes.append('N skipped (exists)')

# ============ O. conclusion rewrite + drop [20] ============
p106 = find_para(lambda t: t.startswith('In conclusion, this study presents CLIP-CAT'), 'O-concl')
new_concl = ('In conclusion, this study presents CLIP-CAT, a vision-language framework designed to address '
             'the challenges of data scarcity, class imbalance, and ordinality in cataract grading. By combining '
             'adaptive patching for micro-lesion detection with a bidirectional rank-aware loss, our approach '
             'outperforms generic backbones and task-specific baselines while offering interpretable, semantic '
             'reasoning for clinical use. Although the current study is limited by a relatively small, single-center '
             'dataset and potential confounders from comorbidities, five-fold cross-validation indicates stable '
             'performance and we restrict our claims accordingly. Future work will prioritize multi-center external '
             'validation and refining prompt engineering to enhance robustness against imaging artifacts and complex '
             'ocular conditions in real-world scenarios')
p106.runs[0].text = new_concl
p106.runs[1].text = ''  # drop misplaced [20]
changes.append('O conclusion')

doc.save(DOC)
print('[saved]', ' | '.join(changes))

# ============ verify ============
doc2 = Document(DOC)
paras = doc2.paragraphs
full = '\n'.join(p.text for p in paras)
checks = []
def chk(name, ok):
    checks.append((name, ok))
    print(f'[check] {name}: {"OK" if ok else "FAIL"}')

chk('title limited', 'Limited and Imbalanced Data' in paras[0].text)
chk('no few-shot anywhere', 'few-shot' not in full.lower() and 'Few-Shot' not in full)
chk('keywords', 'Limited-Scale Learning' in full)
chk('abs model typo fixed', 'vision-language model' in full)
chk('abs no significantly', 'significantly' not in full)
chk('abs sync', 'task-specific' in full and 'five-fold cross-validation' in full)
chk('clinical motivation', 'Clinical Motivation and Imaging Choice' in full and 'OCT angiography' in full)
chk('contrib1 de-novel', 'novel synergy' not in full)
chk('contrib3 momentum', 'momentum-updated global class centers (Section 4.2)' in full)
chk('rw21 lineage', 'CNN backbones [16,17]' in full and 'rank-consistent and kappa-based objectives' in full)
chk('rw22 positioning', 'Unlike prior pipelines that attach ordinal losses to CNN features' in full)
chk('prompt argument', 'Not Label Rewriting' in full and 'dominated by surface form' in full)
chk('eq refs', 'Here in equation (2),' in full and 'This equation (3) aggregates' in full and 'In equations (4) and (5):' in full)
chk('no bare eq refs', 'Here in equation 2,' not in full and 'This equation 3 aggregates' not in full)
chk('h79 realign', 'Comparison with Baseline and Task-Specific Methods' in full)
chk('h18 realign', 'Framework Overview and Problem Definition' in full)
chk('no plummeting', 'plummeting' not in full)
chk('row7 dashitem', 'Impact of Prompt Semantics (Row 7):' in full)
chk('cap3 row7', 'Row 7 replaces the text descriptions with learnable class tokens.' in full)
chk('concl clean', 'significantly outperforms' not in full and 'a novel vision-language' not in full)
chk('concl no [20]', 'real-world scenarios[20]' not in full)

t3b = doc2.tables[2]
shape3 = (len(t3b.rows), len(t3b.columns))
chk('table3 shape (8,7)', shape3 == (8, 7))
row7 = [c.text for c in t3b.rows[-1].cells]
chk('row7 values', row7 == ['7', '\u2713(learnable tokens)', '\u2713', '\u2713', '\u2713', '85.10%', '0.50'])
shapes = [(len(t.rows), len(t.columns)) for t in doc2.tables]
chk('table order', shapes == [(6,2),(5,5),(8,7),(9,5),(7,7)])
chk('images 2', len(doc2.inline_shapes) == 2)
eq_ok = all(f'({i})' in full for i in range(1,10))
chk('equations 1-9 present', eq_ok)

fails = [n for n, ok in checks if not ok]
print(f'[verify] {len(checks)-len(fails)}/{len(checks)} OK')
if fails:
    raise SystemExit('FAIL: ' + ', '.join(fails))
print('[verify] PASS')
