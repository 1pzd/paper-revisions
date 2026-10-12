# -*- coding: utf-8 -*-
"""核实 [14] 引用仍在 + Road.md Phase 6 同步"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)
hits = []
for i, p in enumerate(doc.paragraphs):
    if '[14]' in p.text:
        hits.append((i, p.text[:120]))
print('=== [14] occurrences:', len(hits))
for i, t in hits:
    print(f'  P{i:03d}: {t!r}')

# 引用存在性审计: [1]-[23] 各自在正文中的出现
import re
body = []
for p in doc.paragraphs:
    body.append(p.text)
alltext = '\n'.join(body)
missing = []
for n in range(1, 24):
    if f'[{n}]' not in alltext:
        missing.append(n)
print('=== refs never cited in text:', missing if missing else 'none (all [1]-[23] cited)')

# ---- Road.md 更新 ----
PATH = r'E:\修改自己的论文\Road.md'
with open(PATH, 'r', encoding='utf-8') as f:
    lines = f.readlines()

hits_n = 0
for i, ln in enumerate(lines):
    if ln.startswith('> **最后更新**'):
        lines[i] = '> **最后更新**：2026-10-09（Phase 1 ✅ / Phase 2 ✅ / Phase 3 ✅ / Phase 4 ✅ / Phase 5 ✅ / Phase 6 ✅，Final 进行中）\n'
        hits_n += 1
    elif ln.startswith('| Phase 6 |'):
        cells = ln.rstrip('\n').split('|')
        cells[-2] = ' ✅ '
        lines[i] = '|'.join(cells) + '\n'
        hits_n += 1
    elif ln.startswith('### Phase 6 ⬜'):
        lines[i] = ln.replace('⬜', '✅')
        hits_n += 1

# Phase 6 清单勾选
phase6_done = [
    'Fig. 1 重绘', 'Fig. 2 重绘', 'Fundus image', '空 referenceitem', '图注',
]
for i, ln in enumerate(lines):
    s = ln.strip()
    if s.startswith('- [ ]') and any(k in ln for k in phase6_done):
        lines[i] = ln.replace('- [ ]', '- [x]', 1)
        hits_n += 1

# 更新日志追加
for i, ln in enumerate(lines):
    if ln.startswith('## 5. 更新日志'):
        # 找到日志表末尾（下一个 ## 或文件尾），在其前插入
        j = i + 1
        while j < len(lines) and not lines[j].startswith('## '):
            j += 1
        entry = ('| 2026-10-09 | **Phase 6 完成 ✅**：matplotlib 重绘 Fig.1（仅正文组件：双分支+注意力融合+文本对齐+rank loss，去 Cross-Modal/Report 等）与 Fig.2（单峰分布+双向距离惩罚，(y-c)/(c-y) 标注）、图注同步、Table 1 六行 Fundus→Slit-lamp image、删除末尾 2 空 referenceitem、图片 blob 替换；验证 PASS（134 段/5 表/2 图/公式 1–9/23 条文献） |\n')
        lines.insert(j - 1, entry)
        hits_n += 1
        break

with open(PATH, 'w', encoding='utf-8') as f:
    f.writelines(lines)
print(f'[road] hits={hits_n} (expect 5+)')
print('[saved]', PATH)
