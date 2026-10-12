import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

print('=== Table 1 (prompts) all cells ===')
t0 = doc.tables[0]
for i, row in enumerate(t0.rows):
    cells = [c.text.strip().replace('\n', ' | ') for c in row.cells]
    print(f'r{i}: {cells}')

print()
print('=== Table 4 first column (grade names) ===')
t3 = doc.tables[3]
for i, row in enumerate(t3.rows):
    print(f'r{i}: {row.cells[0].text.strip()} | {row.cells[4].text.strip()}')

print()
print('=== Table 5 header row ===')
t4 = doc.tables[4]
print(' | '.join(c.text.strip() for c in t4.rows[0].cells))

print()
print('=== keyword paragraphs ===')
KEYS = ('NC', 'cortical', 'nuclear', 'C0', 'C1', 'C2', 'LOCS', 'grade', 'dataset',
        'images were', '2,000', '2000', '1673', 'G0', 'G1', 'G2', 'G5', 'consensus',
        'long-tail', 'long-tailed', 'imbalance', 'imbalanced', 'tail', 'head class')
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if any(k.lower() in t.lower() for k in KEYS):
        print(f'P{i:03d} [{p.style.name}] {t[:300]}')
