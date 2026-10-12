# -*- coding: utf-8 -*-
"""Phase 6 取证: Table1 文案 / 图图注互指 / 末尾空段 / 图片 part 信息"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

print('=== paragraphs total:', len(doc.paragraphs))

t1 = doc.tables[0]
print('=== Table1 shape:', len(t1.rows), 'x', len(t1.columns))
for ri, row in enumerate(t1.rows):
    for ci, cell in enumerate(row.cells):
        print(f'  r{ri}c{ci}: {cell.text!r}')

print('=== fig caption / refs in text:')
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if ('Figure' in t or 'Fig.' in t) and t.strip():
        print(f'  P{i:03d} [{p.style.name}] {t[:150]!r}')

print('=== last 8 paragraphs:')
n = len(doc.paragraphs)
for i, p in enumerate(doc.paragraphs[-8:], start=n - 8):
    print(f'  P{i:03d} [{p.style.name}] runs={len(p.runs)} text={p.text[:60]!r}')

refs = [p for p in doc.paragraphs if p.style.name == 'referenceitem']
print('=== referenceitem count:', len(refs))
print('   last 4:')
for p in refs[-4:]:
    print('   ', repr(p.text[:70]))

print('=== inline shapes:', len(doc.inline_shapes))
for i, sh in enumerate(doc.inline_shapes):
    print(f'  shape{i}: type={sh.type} cx={sh.width} cy={sh.height} ({sh.width/914400:.2f}x{sh.height/914400:.2f} in)')
for rid, rel in doc.part.rels.items():
    if "image" in rel.reltype:
        part = rel.target_part
        print(f'  rel {rid}: {part.partname} ct={part.content_type} len={len(part.blob)}')

# heading outline to anchor Phase 6
print('=== headings:')
for i, p in enumerate(doc.paragraphs):
    if p.style.name in ('heading1', 'heading2', 'Title'):
        print(f'  P{i:03d} [{p.style.name}] {p.text[:80]!r}')
