import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

def replace_in_para(p, old, new):
    runs = p.runs
    texts = [r.text for r in runs]
    joined = ''.join(texts)
    if old in joined:
        for r in runs:
            if old in r.text:
                r.text = r.text.replace(old, new)
                return True
        idx = joined.index(old); end = idx + len(old); pos = 0; first = True
        for r in runs:
            L = len(r.text); s, e = pos, pos + L
            if s < end and e > idx:
                if first:
                    pre = r.text[:max(0, idx - s)]
                    post = r.text[max(0, min(L, end - s)):]
                    r.text = pre + new + post
                    first = False
                else:
                    r.text = '' if e <= end else r.text[end - s:]
            pos = e
        return True
    return False

# 1. P099 text 85.9 -> 85.8
hit = False
for p in doc.paragraphs:
    if p.text.startswith('Per-Grade Observations.'):
        hit = replace_in_para(p, '85.9%', '85.8%')
        break
print('P099 85.9->85.8:', hit)

# 2. Table4 row G2 (index 3) F1 cell 85.9 -> 85.8
t4 = doc.tables[3]
cell = t4.rows[3].cells[3]
old = cell.text.strip()
cell.paragraphs[0].runs[0].text = '85.8'
for r in cell.paragraphs[0].runs[1:]:
    r._r.getparent().remove(r._r)
print('Table4 G2 F1:', old, '->', cell.text.strip())

doc.save(DOC)

# verify
doc2 = Document(DOC)
body = [[int(x) for x in r.cells[1:]] for r in [doc2.tables[4].rows[i] for i in range(1, 7)]]
mat = [[c.text.strip() for c in r.cells] for r in doc2.tables[4].rows]
body = [[int(x) for x in row[1:]] for row in mat[1:]]
m4 = [[c.text.strip() for c in r.cells] for r in doc2.tables[3].rows]
ok = True
for i in range(6):
    tp = body[i][i]
    col = sum(body[k][i] for k in range(6))
    row = sum(body[i])
    prec = tp / col * 100 if col else 0
    rec = tp / row * 100
    f1 = 2 * prec * rec / (prec + rec) if (prec + rec) else 0
    got = (round(prec, 1), round(rec, 1), round(f1, 1))
    want = tuple(float(x) for x in m4[i + 1][1:4])
    if got != want:
        ok = False
        print(f'row{i}: matrix{got} vs table{want}')
print('[verify] Table4 metrics match matrix:', ok)

# macro/weighted recompute
P = [float(m4[i+1][1]) for i in range(6)]
R = [float(m4[i+1][2]) for i in range(6)]
F = [float(m4[i+1][3]) for i in range(6)]
sup = [int(m4[i+1][4]) for i in range(6)]
macro = (round(sum(P)/6,1), round(sum(R)/6,1), round(sum(F)/6,1))
w = lambda v: round(sum(s*x for s, x in zip(sup, v))/sum(sup), 1)
wavg = (w(P), w(R), w(F))
print('[verify] macro:', macro, 'table:', tuple(float(x) for x in m4[7][1:4]))
print('[verify] weighted:', wavg, 'table:', tuple(float(x) for x in m4[8][1:4]))
alltext = '\n'.join(p.text for p in doc2.paragraphs)
print('[verify] P099 has 85.8%:', '85.8%' in alltext and '85.9%' not in alltext)
print('[verify] drawings:', len(doc2.inline_shapes), 'tables:', len(doc2.tables))
print('RESULT:', 'PASS' if ok and macro == tuple(float(x) for x in m4[7][1:4]) and wavg == tuple(float(x) for x in m4[8][1:4]) else 'CHECK')
