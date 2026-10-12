# -*- coding: utf-8 -*-
"""Phase 6 取证3: 图注 run 结构 + 全文 Fundus 残留扫描"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

for i in (29, 65):
    p = doc.paragraphs[i]
    print(f'--- P{i:03d} [{p.style.name}] runs={len(p.runs)}')
    for j, r in enumerate(p.runs):
        print(f'  run{j}: {r.text!r} bold={r.bold}')

print('=== Fundus/fundus scan ===')
for i, p in enumerate(doc.paragraphs):
    if 'fundus' in p.text.lower():
        print(f'  para P{i:03d}: {p.text[:110]!r}')
for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if 'fundus' in cell.text.lower():
                print(f'  table{ti} r{ri}c{ci}: {cell.text[:80]!r}')
print('=== empty referenceitem ===')
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'referenceitem' and not p.text.strip():
        print(f'  P{i:03d} runs={len(p.runs)}')
