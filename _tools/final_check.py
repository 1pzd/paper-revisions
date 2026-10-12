# -*- coding: utf-8 -*-
"""终检: 引用顺序+公式完整性+图表+关键文本 一次性核验"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn as Q

DOC = r'E:\修改自己的论文\CLIP.docx'
CIT = re.compile(r'\[(\d+(?:\s*[-,]\s*\d+)*)\]')

def parse_ordered(b):
    out = []
    for part in b.split(','):
        part = part.strip()
        if '-' in part:
            a, c = part.split('-')
            out.extend(range(int(a.strip()), int(c.strip()) + 1))
        elif part:
            out.append(int(part))
    return out

doc = Document(DOC)
ps = doc.paragraphs
ok = True
def chk(cond, msg):
    global ok
    print(('  OK ' if cond else '  FAIL ') + msg)
    if not cond:
        ok = False

ref_start = next(i for i, p in enumerate(ps) if p.text.strip() == 'References')
body = ps[:ref_start]

seq = []
for p in body:
    for m in CIT.finditer(p.text):
        for n in parse_ordered(m.group(1)):
            if n not in seq:
                seq.append(n)
chk(seq == list(range(1, 24)), f'cite order 1..23: {seq}')

ref_paras = [p for p in ps if p.style.name == 'referenceitem' and p.text.strip()]
chk(len(ref_paras) == 23 and not any(p.style.name == 'referenceitem' and not p.text.strip() for p in ps), '23 refs, no empty tail')
chk('de la Torre' in ref_paras[13].text and 'Cao, W.' in ref_paras[14].text and 'Radford' in ref_paras[16].text, '[14]Torr [15]Cao [17]Radford')

alltext = '\n'.join(p.text for p in body)
chk('11-13' in alltext and '[5,6,11-13,15,17]' in alltext, 'Table2 cite bracket fully collapsed')
chk('CLIP [17]' in alltext and 'augmentation[23]' in alltext and 'CORAL ordinal objective [15]' in alltext, 'key cites intact')

shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
chk(shapes == [(6, 2), (10, 5), (8, 7), (9, 5), (7, 7)], f'tables: {shapes}')
chk(len(list(doc.element.body.iter(Q('w:drawing')))) == 2, '2 drawings')

# 9 个公式段各含 oMath 且标签 (1)-(9)
eqp = [p for p in ps if 'equation' in (p.style.name or '').lower()]
labels, withmath = [], 0
for p in eqp:
    nodes = [t.text or '' for t in p._p.iter() if t.tag in (Q('m:t'), Q('w:t'))]
    joined = ''.join(nodes)
    m = re.search(r'\((\d)\)\s*$', joined.strip())
    if m:
        labels.append(m.group(1))
    if list(p._p.iter(Q('m:oMath'))):
        withmath += 1
chk(labels == [str(i) for i in range(1, 10)], f'eq labels: {labels}')
chk(withmath == len(eqp) and len(eqp) >= 9, f'every eq para has oMath ({withmath}/{len(eqp)})')

low = alltext.lower()
chk('few-shot' not in low and 'significantly outperforms' not in low and 'plummeting' not in low, 'no promo residue')
chk('Limited and Imbalanced Data' in ps[0].text, 'title updated')
chk('Slit-lamp image' in doc.tables[0].cell(0, 1).text, 'Table1 Slit-lamp')
chk('Per-Grade Performance' in alltext and 'Generalizability and Limitations' in alltext, 'new sections present')

print('RESULT:', 'PASS' if ok else 'FAIL')
print(f'[info] paragraphs={len(ps)} tables={len(doc.tables)} refs={len(ref_paras)}')
