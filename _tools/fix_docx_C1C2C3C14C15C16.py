# -*- coding: utf-8 -*-
"""Docx-only continuation of fix_C1C2C3C14C15C16.py (txt already applied + verified).

Applies the C1/C2/C3/C14/C15/C16 edits to CLIP.docx, including the
reconstruction of mangled performance-table rows 05-09.
Backups were already created by the previous run (pre-edit state).
"""
import sys
from docx import Document

ROOT = r"D:\修改论文"
DOCX = ROOT + r"\CLIP.docx"


def die(msg):
    sys.stderr.write("FAIL: %s\n" % msg)
    sys.exit(1)


doc = Document(DOCX)


def replace_in_para(p, old, new, label):
    runs = p.runs
    texts = [r.text for r in runs]
    full = "".join(texts)
    idx = full.find(old)
    if idx == -1:
        die("docx %s: anchor not found" % label)
    if full.find(old, idx + 1) != -1:
        die("docx %s: anchor ambiguous" % label)
    end = idx + len(old)
    cum = []
    pos = 0
    for tx in texts:
        cum.append(pos)
        pos += len(tx)
    new_texts = list(texts)
    for i, tx in enumerate(texts):
        rs, re = cum[i], cum[i] + len(tx)
        if re <= idx or rs >= end:
            continue
        if rs <= idx < re:
            head = tx[: idx - rs]
            tail = tx[end - rs :] if end <= re else ""
            new_texts[i] = head + new + tail
        elif rs < end <= re:
            new_texts[i] = tx[end - rs :]
        else:
            new_texts[i] = ""
    for r, nt in zip(runs, new_texts):
        r.text = nt


def para_by_start(prefix):
    hits = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if len(hits) != 1:
        die("docx para %r: found %d" % (prefix[:40], len(hits)))
    return hits[0]


# C3 + C2
p074 = para_by_start("To validate our framework")
replace_in_para(
    p074,
    "of 1,673 slit-lamp images, with ground truth established by three senior "
    "ophthalmologists via consensus voting.",
    "of 1,673 slit-lamp images. The cohort included patients across all age stages "
    "and of both sexes. All images were collected via routine clinical slit-lamp "
    "photography; the specific slit-lamp models and acquisition parameters were not "
    "recorded. Ground truth was established by three senior ophthalmologists via "
    "consensus voting: each image was graded by all three graders, and the final "
    "label was determined by consensus. As the reference standard reflects consensus "
    "rather than independent parallel ratings, inter-rater \u03ba is not reported.",
    "C3+C2 dataset desc",
)

# C1 - bootstrap CI
p083 = para_by_start("Performance Superiority:")
replace_in_para(
    p083,
    "by 4.21%. While",
    "by 4.21%. A nonparametric bootstrap 95% confidence interval for CLIP-CAT "
    "accuracy (10,000 resamples of the 335 test images) is [83.3%, 90.4%]. While",
    "C1 bootstrap CI",
)

# C14
c14_hits = [p for p in doc.paragraphs if "carry wide confidence intervals;" in p.text]
if len(c14_hits) != 1:
    die("docx C14 anchor: found %d" % len(c14_hits))
replace_in_para(
    c14_hits[0],
    "carry wide confidence intervals;",
    "carry wide confidence intervals (exact binomial 95% CI for recall: "
    "G4 75.0% [19.4%\u201399.4%]; G5 66.7% [9.4%\u201399.2%]);",
    "C14 G4G5 CI",
)

# C16
c16_hits = [p for p in doc.paragraphs if "G2\u2194G3, 13 cases" in p.text]
if len(c16_hits) != 1:
    die("docx C16 anchor: found %d" % len(c16_hits))
replace_in_para(c16_hits[0], "G2\u2194G3, 13 cases", "G2\u2194G3, 13 cases: 10 + 3", "C16 bidi")

# C15 - captions
p086 = para_by_start("Table 2. Performance comparison")
replace_in_para(
    p086,
    "Table 2. Performance comparison with baseline and task-specific models.",
    "Table 2. Performance comparison with baseline and task-specific models. "
    "Values are means over 3 independent runs (seeds 42, 123, and 2024); "
    "values for CLIP-CAT are reported as mean \u00b1 standard deviation.",
    "C15 tbl2 caption",
)
p100 = para_by_start("Table 4. Per-grade precision, recall, and F1 of")
replace_in_para(
    p100,
    "Table 4. Per-grade precision, recall, and F1 of CLIP-CAT on the held-out test subset.",
    "Table 4. Per-grade precision, recall, and F1 of CLIP-CAT on the held-out test subset. "
    "Single-run results (seed 42).",
    "C15 tbl4 caption",
)
p102 = para_by_start("Table 5. Confusion matrix")
replace_in_para(
    p102,
    "Table 5. Confusion matrix on the held-out test set (335 images). "
    "Rows represent true grades, columns represent predicted grades.",
    "Table 5. Confusion matrix on the held-out test set (335 images). "
    "Rows represent true grades, columns represent predicted grades. "
    "Single-run results (seed 42).",
    "C15 tbl5 caption",
)


# C1 - rebuild mangled rows 05-09 of performance table (tables[1])
def set_cell(cell, text):
    paras = cell.paragraphs
    p = paras[0]
    for r in p.runs:
        r.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)


t1 = doc.tables[1]
new_rows = [
    ("CLIP (linear probe)", "CLIP ViT-L/14", "79.50%", "1.05", "0.73"),
    ("ResNet50 + CORAL", "CNN + Ordinal", "80.30%", "0.90", "0.75"),
    ("CDRF-Net", "Task-specific", "83.10%", "0.72", "0.81"),
    ("LOCS III hybrid", "Task-specific", "82.65%", "0.78", "0.80"),
    ("CLIP-CAT (ours)", "Vision-Language", "86.76% \u00b1 0.20%", "0.42 \u00b1 0.02", "0.88 \u00b1 0.01"),
]
for ri, rowvals in enumerate(new_rows, start=5):
    cells = t1.rows[ri].cells
    if len(cells) != 5:
        die("docx tbl2 row %d: %d cells" % (ri, len(cells)))
    for ci, val in enumerate(rowvals):
        set_cell(cells[ci], val)

# ---- verification
paras = doc.paragraphs
fulltext = "\n".join(p.text for p in paras)
if "83.3%, 90.4%" not in fulltext:
    die("docx: C1 bootstrap CI missing")
if "exact binomial 95% CI for recall" not in fulltext:
    die("docx: C14 CI missing")
if "G2\u2194G3, 13 cases: 10 + 3" not in fulltext:
    die("docx: C16 wording missing")
if fulltext.count("Single-run results (seed 42).") != 2:
    die("docx: C15 single-run marks != 2")
if "Values are means over 3 independent runs" not in fulltext:
    die("docx: C15 3-run mark missing")
if "inter-rater \u03ba is not reported" not in fulltext:
    die("docx: C2 kappa statement missing")
if "slit-lamp models and acquisition parameters were not recorded" not in fulltext:
    die("docx: C3 device note missing")
if "across all age stages and of both sexes" not in fulltext:
    die("docx: C3 demographics missing")

t1texts = [[c.text for c in r.cells] for r in t1.rows]
if t1texts[9] != ["CLIP-CAT (ours)", "Vision-Language", "86.76% \u00b1 0.20%", "0.42 \u00b1 0.02", "0.88 \u00b1 0.01"]:
    die("docx tbl2 CLIP-CAT row wrong: %r" % t1texts[9])
if t1texts[5][0] != "CLIP (linear probe)" or t1texts[8][0] != "LOCS III hybrid":
    die("docx tbl2 rows 5/8 wrong")
if "zero-shot" not in t1texts[4][0]:
    die("docx tbl2 zero-shot row wrong: %r" % t1texts[4][0])

import re

n_tbl = len(doc.tables)
n_om = len(re.findall(r"<m:oMath>", doc.element.body.xml))
n_ompara = len(re.findall(r"<m:oMathPara>", doc.element.body.xml))
n_ref = sum(1 for p in paras if "referenceitem" in p.style.name.lower())
n_img = sum(1 for p in paras if "blip" in p._p.xml)
print("DOCX paras=%d tables=%d oMath=%d oMathPara=%d refs=%d imgs=%d"
      % (len(paras), n_tbl, n_om, n_ompara, n_ref, n_img))
if n_tbl != 5:
    die("docx tables != 5")
if n_om != 45 or n_ompara != 8:
    die("docx oMath changed")
if n_ref != 26:
    die("docx refs != 26")
if n_img != 4:
    die("docx imgs != 4")

doc.save(DOCX)
print("DOCX OK")
