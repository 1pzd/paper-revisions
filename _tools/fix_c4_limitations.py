# -*- coding: utf-8 -*-
"""C4 honest fix (refuses fabricated baselines): add a fourth, explicit limitation
explaining why medical VLM baselines (BiomedCLIP/MedCLIP/PubMedCLIP) are not compared,
to CLIP-revised.txt §4.4 and CLIP.docx P109. No numbers are invented."""
import io
import shutil
import sys
from pathlib import Path

from docx import Document

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

ROOT = Path(r"D:\修改论文")
TXT = ROOT / "CLIP-revised.txt"
DOCX = ROOT / "CLIP.docx"
BAK = ROOT / "_backup"

FOURTH = (
    "Fourth, the baseline comparison does not include medical vision-language models such as "
    "BiomedCLIP, MedCLIP, and PubMedCLIP. These models are pretrained on radiology, pathology, "
    "or biomedical figure-caption corpora rather than slit-lamp anterior-segment photography, "
    "so zero-shot transfer to LOCS III grading would conflate imaging-modality mismatch with "
    "representation quality. A fair comparison would require fine-tuning each backbone under "
    "the same long-tailed protocol, which is beyond the scope of the present limited-data "
    "study; our zero-shot CLIP and linear-probe baselines instead isolate the contribution of "
    "vision-language pretraining and supervised adaptation within the target modality. We "
    "leave the comparison with domain-specific medical vision-language backbones to future "
    "multi-center work."
)

OLD_LEAD = "This study has three concrete limitations."
NEW_LEAD = "This study has four concrete limitations."
OLD_TAIL = "and external multi-center validation remains outstanding."
NEW_TAIL = OLD_TAIL + " " + FOURTH


def replace_span_in_para(p, old, new):
    """Run-aware replacement preserving run formatting."""
    runs = list(p.runs)
    joined = "".join(r.text for r in runs)
    i = joined.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos = 0
    placed = False
    for r in runs:
        t = r.text
        s, e = pos, pos + len(t)
        pos = e
        if e <= i or s >= j:
            continue
        left = t[: i - s] if i > s else ""
        right = t[j - s :] if s < j <= e else ""
        if not placed:
            r.text = left + new + right
            placed = True
        else:
            r.text = right
    return True


def main():
    # 1. backups
    shutil.copy2(TXT, BAK / "CLIP-revised_pre_C4.txt")
    shutil.copy2(DOCX, BAK / "CLIP_pre_C4.docx")
    print("[backup] CLIP-revised_pre_C4.txt / CLIP_pre_C4.docx written")

    # 2. txt
    text = TXT.read_text(encoding="utf-8")
    assert text.count(OLD_LEAD) == 1, f"txt lead count={text.count(OLD_LEAD)}"
    assert text.count(OLD_TAIL) == 1, f"txt tail count={text.count(OLD_TAIL)}"
    text = text.replace(OLD_LEAD, NEW_LEAD)
    text = text.replace(OLD_TAIL, NEW_TAIL)
    TXT.write_text(text, encoding="utf-8")
    print("[txt] patched OK")

    # 3. docx
    doc = Document(DOCX)
    hits = 0
    for p in doc.paragraphs:
        if OLD_LEAD in p.text:
            hits += 1
            print("[docx] target paragraph runs:")
            for k, r in enumerate(p.runs):
                print(f"  run{k}: bold={r.bold} italic={r.italic} text={r.text[:60]!r}")
            ok1 = replace_span_in_para(p, OLD_LEAD, NEW_LEAD)
            ok2 = replace_span_in_para(p, OLD_TAIL, NEW_TAIL)
            assert ok1 and ok2, (ok1, ok2)
            print("[docx] paragraph patched")
    assert hits == 1, f"docx target hits={hits}"
    doc.save(DOCX)
    print("[docx] saved OK")

    # 4. verification
    text2 = TXT.read_text(encoding="utf-8")
    assert NEW_LEAD in text2 and OLD_LEAD not in text2
    assert FOURTH in text2
    doc2 = Document(DOCX)
    found = [p.text for p in doc2.paragraphs if "four concrete limitations" in p.text]
    assert len(found) == 1, len(found)
    assert found[0].count("Fourth, the baseline comparison does not include medical") == 1
    assert "three concrete limitations" not in found[0]
    assert found[0].startswith("Limitations."), found[0][:40]
    print("[verify] txt + docx both contain the fourth limitation, lead-in intact")
    print("[verify] docx paragraph head:", found[0][:120])
    print("[verify] docx paragraph tail:", found[0][-120:])


if __name__ == "__main__":
    main()
