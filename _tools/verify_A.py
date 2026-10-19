# -*- coding: utf-8 -*-
"""Verify A1-A4 edits in CLIP.docx: show edited passages, scan tables for weights, list [6] citations."""
from docx import Document

doc = Document('../CLIP.docx')
print('=== edited-passage spot check ===')
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if ('rounded up to the nearest integer' in t or 'two-fold' in t
            or '1,171 training' in t or 'n_0=229' in t or 'Prototype Smoothing' in t):
        print('P%03d: %s' % (i, t[:420]))
        print()

print('=== weight numbers anywhere (paras + tables) ===')
keys = ['w_0', '27.9', '37.2', 'w_1=1.2', 'w_3=5.5', 'n_0=229', '13.94']
for i, p in enumerate(doc.paragraphs):
    if any(k in p.text for k in keys):
        print('PARA P%03d: %s' % (i, p.text[:300]))
for ti, tbl in enumerate(doc.tables):
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            if any(k in cell.text for k in keys):
                print('TBL T%d R%d C%d: %s' % (ti, ri, ci, cell.text[:200]))
print('(end weight scan)')

print('=== [6] citation scan (docx bracket style) ===')
import re
for i, p in enumerate(doc.paragraphs):
    for m in re.finditer(r'\[\s*6[\s,\]]|\[\s*[\d,\s]*,\s*6\s*[,\]]|[\d\s]-\s*6\s*\]', p.text):
        print('P%03d %r' % (i, p.text[max(0, m.start()-60):m.end()+10]))
print('(end [6] scan)')
print('=== done ===')
