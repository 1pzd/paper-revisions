# -*- coding: utf-8 -*-
"""Dump Phase-2 context: grade names, long-tail text, tables, body order."""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

DOC = r"E:\修改自己的论文\CLIP.docx"
doc = Document(DOC)

print("=== KEY PARAGRAPHS (full text) ===")
for i in [75, 80, 81, 82, 85, 86, 87, 88, 89, 90]:
    print(f"--- P{i:03d} ---")
    print(doc.paragraphs[i].text)

print()
print("=== TABLE 0 (Table 1 in paper: grade descriptions) ===")
t = doc.tables[0]
for r, row in enumerate(t.rows):
    print(f"row{r}: " + " || ".join(c.text.replace(chr(10), ' ') for c in row.cells))

print()
print("=== TABLE 1 (Table 2 in paper: comparison) ===")
t = doc.tables[1]
for r, row in enumerate(t.rows):
    print(f"row{r}: " + " || ".join(c.text.replace(chr(10), ' ') for c in row.cells))

print()
print("=== TABLE 2 (Table 3 in paper: ablation) ===")
t = doc.tables[2]
for r, row in enumerate(t.rows):
    print(f"row{r}: " + " || ".join(repr(c.text) for c in row.cells))

print()
print("=== BODY ORDER around tables (para 79..93) ===")
body = doc.element.body
from docx.text.paragraph import Paragraph
from docx.table import Table
pi = 0
ti = 0
for child in body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        p = Paragraph(child, doc)
        if 79 <= pi <= 93:
            style = p.style.name if p.style is not None else '?'
            print(f"[P{pi:03d}]({style}) {p.text[:60]}")
        pi += 1
    elif tag == 'tbl':
        print(f"  <TABLE {ti}> after P{pi-1:03d}")
        ti += 1

print()
print("=== ABLATION TABLE XML (style/borders) ===")
t = doc.tables[2]
print(t._tbl.xml[:3000])
