# -*- coding: utf-8 -*-
"""Final: 重编号后纠正版验证（期望映射按真实文本首现序；图像/公式用稳健计法）"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn as Q

DOC = r'E:\修改自己的论文\CLIP.docx'
CIT = re.compile(r'\[(\d+(?:\s*[-,]\s*\d+)*)\]')

def parse_ordered(body):
    out = []
    for part in body.split(','):
        part = part.strip()
        if '-' in part:
            a, b = part.split('-')
            out.extend(range(int(a.strip()), int(b.strip()) + 1))
        elif part:
            out.append(int(part))
    return out

doc = Document(DOC)
ps = doc.paragraphs
ok = True
def chk(cond, msg):
    global ok
    print(('  OK ' if cond else '  FAIL ') + msg)
    if not cond:
        ok = False

ref_start = next(i for i, p in enumerate(ps) if p.text.strip() == 'References')
body = ps[:ref_start]

# 1. 首现顺序
seq = []
for p in body:
    for m in CIT.finditer(p.text):
        for n in parse_ordered(m.group(1)):
            if n not in seq:
                seq.append(n)
chk(seq == list(range(1, 24)), f'first-occurrence order 1..23: {seq}')

# 2. 文献表定位（真实映射: de la Torre=14, Cao=15, Radford=17, Lin=21）
ref_paras = [p for p in ps if p.style.name == 'referenceitem' and p.text.strip()]
chk(len(ref_paras) == 23, f'23 non-empty refs (got {len(ref_paras)})')
empty_refs = [p for p in ps if p.style.name == 'referenceitem' and not p.text.strip()]
chk(len(empty_refs) == 0, f'no trailing empty referenceitem (got {len(empty_refs)})')
chk('de la Torre' in ref_paras[13].text, '[14] = de la Torre')
chk('Cao, W.' in ref_paras[14].text, '[15] = Cao CORAL')
chk('Radford' in ref_paras[16].text, '[17] = Radford CLIP')
chk('Lin' in ref_paras[20].text and 'Focal' in ref_paras[20].text, '[21] = Lin Focal loss')
chk('Complementary discriminative region' in ref_paras[4].text, '[5] = CDRF-Net')
chk('LOCS III-based hybrid' in ref_paras[5].text, '[6] = LOCS III hybrid')

# 3. 正文关键引用
alltext = '\n'.join(p.text for p in body)
chk('CLIP [17]' in alltext, "P017 CLIP cite = [17]")
chk('augmentation[23]' in alltext, 'conclusion mixup = [23]')
chk('CORAL ordinal objective [15]' in alltext, 'CORAL cite = [15]')
chk('cross-grade errors [14,15]' in alltext, 'rank/kappa objectives cite [14,15]')
chk('objective [5,15]' in alltext or '[5,15]' in alltext, 'CNN-ordinal coupling cite [5,15]')
chk('[4-6,16]' in alltext, "RW2.1 tail bracket [7-10] -> [4-6,16]")
chk('[11-13]' in alltext, 'backbone cites collapsed [11-13]')

# 4. 表/图稳健计数
shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
chk(shapes == [(6, 2), (10, 5), (8, 7), (9, 5), (7, 7)], f'table shapes: {shapes}')
n_draw = len(list(doc.element.body.iter(Q('w:drawing'))))
n_anchor = len(list(doc.element.body.iter(Q('wp:anchor'))))
chk(n_draw == 2, f'w:drawing count = {n_draw} (inline_shapes={len(doc.inline_shapes)}, anchors={n_anchor})')

# 5. 公式标签（拼接段内 w:t+m:t 顺序文本，容忍跨节点拆分）
eqp = [p for p in ps if 'equation' in (p.style.name or '').lower()]
labels = []
for p in eqp:
    nodes = [t.text or '' for t in p._p.iter() if t.tag in (Q('m:t'), Q('w:t'))]
    joined = ''.join(nodes)
    m = re.search(r'\((\d)\)\s*$', joined.strip())
    if m:
        labels.append(m.group(1))
chk(labels == [str(i) for i in range(1, 10)], f'eq labels (1)-(9): {labels}')

# 6. 全文无残留旧式未压缩引用格式异常 & 无 few-shot/宣传语
low = alltext.lower()
chk('few-shot' not in low and 'significantly outperforms' not in low, 'no few-shot / promo residue')
chk(alltext.count('[21]') >= 1 and 'exp(-' not in alltext, 'focal cite present')

print('RESULT:', 'PASS' if ok else 'FAIL')
print(f'[info] paragraphs={len(ps)} tables={len(doc.tables)} drawings={n_draw}')
