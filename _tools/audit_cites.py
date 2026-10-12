# -*- coding: utf-8 -*-
"""Final: 精确引用审计（支持 [a,b-c,d] 范围/列表解析）：孤儿引用 + 首现顺序 vs 编号顺序"""
import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

CIT = re.compile(r'\[(\d+(?:\s*[-,]\s*\d+)*)\]')

def parse_cites(s):
    out = set()
    for m in CIT.finditer(s):
        body = m.group(1)
        # split by comma
        for part in body.split(','):
            part = part.strip()
            if '-' in part:
                a, b = part.split('-')
                a, b = int(a.strip()), int(b.strip())
                out.update(range(a, b + 1))
            elif part:
                out.add(int(part))
    return out

# 正文 = 参考文献标题之前的段落
first_occ = {}   # n -> 段落索引
all_cited = set()
ref_start = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == 'References':
        ref_start = i
        break
body_paras = doc.paragraphs[:ref_start]

for i, p in enumerate(body_paras):
    for n in parse_cites(p.text):
        all_cited.add(n)
        if n not in first_occ:
            first_occ[n] = i

print('=== 正文引用到的编号:', sorted(all_cited))
n_refs = len([p for p in doc.paragraphs if p.style.name == 'referenceitem' and p.text.strip()])
print('=== 参考文献条数:', n_refs)
orphans = [n for n in range(1, n_refs + 1) if n not in all_cited]
print('=== 孤儿文献(正文从未引用):', orphans if orphans else '无')
over = [n for n in sorted(all_cited) if n > n_refs]
print('=== 引用超出文献表范围:', over if over else '无')

print('=== 首现顺序（按段落先后）:')
order = sorted(first_occ.items(), key=lambda kv: (kv[1], kv[0]))
for n, i in order:
    txt = body_paras[i].text
    label = txt[:60].replace('\n', ' ')
    print(f'  [{n:2d}] first@P{i:03d}: {label!r}')

seq = [n for n, _ in order]
viol = [(seq[k], seq[k + 1]) for k in range(len(seq) - 1) if seq[k] > seq[k + 1]]
print('=== 顺序违例 (前一个编号 > 后一个编号):')
for a, b in viol:
    print(f'  [{a}] 出现在 [{b}] 之前')
if not viol:
    print('  无 —— 编号顺序与首现顺序一致')
print('=== 首现序列:', seq)
