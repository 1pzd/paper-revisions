# -*- coding: utf-8 -*-
"""Phase 6: 图表重绘替换 + Table1 文案 + 图注同步 + 删空参考文献段（幂等守卫: 'Slit-lamp image' in table0 r0c1）"""
import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document
from docx.oxml.ns import qn

DOC = r'E:\修改自己的论文\CLIP.docx'
NEW1 = r'E:\修改自己的论文\new_fig1.png'
NEW2 = r'E:\修改自己的论文\new_fig2.png'

doc = Document(DOC)
ps = doc.paragraphs
t1 = doc.tables[0]

# ---------- idempotency ----------
if 'Slit-lamp image' in t1.rows[0].cells[1].text:
    print('[skip] already applied')
    sys.exit(0)

def set_cell(cell, text):
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p.add_run(text)

# ---------- 1. Table 1: Fundus image -> Slit-lamp image ----------
n_fix = 0
for row in t1.rows:
    c = row.cells[1]
    if 'Fundus image' in c.text:
        set_cell(c, c.text.replace('Fundus image', 'Slit-lamp image'))
        n_fix += 1
print(f'[1] table1 cells fixed: {n_fix}')

# ---------- 2. captions sync ----------
cap1 = ps[29]
assert cap1.style.name == 'figurecaption' and cap1.runs[7].text.startswith(' Overview'), cap1.runs[7].text
cap1.runs[7].text = (' Overview of the proposed CLIP-CAT framework: the global CLIP encoder and '
                     'the local adaptive patching branch are fused via the attention mechanism, '
                     'aligned with grade-specific medical prompts, and trained with the '
                     'bidirectional rank-aware loss.')
print('[2] Fig.1 caption updated')

cap2 = ps[65]
assert cap2.style.name == 'figurecaption' and 'distance penalty mechanism' in cap2.runs[7].text
cap2.runs[7].text = cap2.runs[7].text.replace('the distance penalty mechanism',
                                              'the bidirectional distance penalty mechanism')
print('[2] Fig.2 caption updated')

# ---------- 3. delete empty referenceitem ----------
deleted = 0
for p in list(doc.paragraphs):
    if p.style.name == 'referenceitem' and not p.text.strip() and len(p.runs) == 0:
        p._p.getparent().remove(p._p)
        deleted += 1
print(f'[3] empty referenceitem deleted: {deleted}')

# ---------- 4. replace figure blobs ----------
shapes = doc.inline_shapes
assert len(shapes) == 2, len(shapes)
replaced = []
for shape, path in zip(shapes, (NEW1, NEW2)):
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    rId = blip.get(qn('r:embed'))
    part = doc.part.related_parts[rId]
    with open(path, 'rb') as f:
        data = f.read()
    part._blob = data
    replaced.append((rId, str(part.partname), len(data)))
print('[4] image blobs replaced:', replaced)

doc.save(DOC)
print('[saved]', DOC)

# ---------- verify ----------
doc = Document(DOC)
ps = doc.paragraphs
t1 = doc.tables[0]
ok = True
def chk(cond, msg):
    global ok
    print(('  OK ' if cond else '  FAIL ') + msg)
    if not cond:
        ok = False

chk('Fundus image' not in '|'.join(c.text for r in t1.rows for c in r.cells), 'table1 no Fundus image')
chk(all('Slit-lamp image' in t1.rows[i].cells[1].text for i in range(6)), 'table1 6 rows Slit-lamp')
chk('bidirectional rank-aware loss' in ps[29].text, 'Fig.1 caption mentions bidirectional rank-aware loss')
chk('bidirectional distance penalty' in ps[65].text, 'Fig.2 caption mentions bidirectional')
chk(not [p for p in ps if p.style.name == 'referenceitem' and not p.text.strip()], 'no empty referenceitem')
chk(len(doc.inline_shapes) == 2, '2 inline shapes')
chk(len(doc.tables) == 5, '5 tables')
sizes = []
for rid, rel in doc.part.rels.items():
    if 'image' in rel.reltype:
        sizes.append((str(rel.target_part.partname), len(rel.target_part.blob)))
chk(all(n > 100000 for _, n in sizes), f'image blobs swapped: {sizes}')
chk(len(ps) == 136 - 2, f'paragraphs={len(ps)} (134 expected)')
# referenceitem count
chk(len([p for p in ps if p.style.name == 'referenceitem']) == 23, '23 referenceitem')
print('RESULT:', 'PASS' if ok else 'FAIL')
