import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

# 1) every paragraph whose text contains a (N) label pattern
print('===== paragraphs containing (N) pattern =====')
for i, p in enumerate(doc.paragraphs):
    if re.search(r'\(\d\)', p.text):
        print(f'P{i} [{p.style.name}]: {p.text!r}')

# 2) all text nodes (w:t + m:t) in every equation-styled paragraph
print()
print('===== equation paragraphs: w:t and m:t nodes =====')
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'equation':
        xml = p._p.xml
        wt = re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml)
        mt = re.findall(r'<m:t[^>]*>([^<]*)</m:t>', xml)
        print(f'P{i}: w:t={wt!r}')
        print(f'     m:t={mt!r}')

# 3) raw XML of P37 (suspected duplicate label)
print()
print('===== P37 raw xml (truncated 3500) =====')
print(doc.paragraphs[37]._p.xml[:3500])

# 4) comparison: original backup, paragraph with same duplicate structure
print()
print('===== backup original: search all (N)-like labels in w:t =====')
bak = Document(r'E:\修改自己的论文\_backup\CLIP_original.docx')
for i, p in enumerate(bak.paragraphs):
    xml = p._p.xml
    wt = re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml)
    joined = ''.join(wt)
    if re.search(r'\(\s*\d\s*\)', joined) or 'SEQ' in joined:
        print(f'bak P{i} [{p.style.name}]: w:t={wt!r}')
