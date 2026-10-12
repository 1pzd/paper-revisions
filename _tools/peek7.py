import sys, io, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
from docx import Document

DOC = r'E:\修改自己的论文\CLIP.docx'
doc = Document(DOC)

# ---- Table 3 structure ----
t3 = doc.tables[2]
print('===== Table 3 (tables[2]) =====')
for ri, row in enumerate(t3.rows):
    cells = [c.text for c in row.cells]
    print(f'row{ri}: {cells}')

# ---- citation map: which paragraphs cite which numbers ----
print()
print('===== citation map =====')
for i, p in enumerate(doc.paragraphs):
    for m in re.finditer(r'\[(\d+(?:[,\-\u2013]\d+)*)\]', p.text):
        print(f'P{i} [{p.style.name}] cites {m.group(0)} :: {p.text[max(0,m.start()-40):m.end()+15]!r}')

# ---- 'equation N' mentions ----
print()
print('===== equation mentions =====')
for i, p in enumerate(doc.paragraphs):
    for m in re.finditer(r'[Ee]quation\s*\d+', p.text):
        print(f'P{i}: {m.group(0)!r} :: {p.text[max(0,m.start()-45):m.end()+25]!r}')

# ---- promotional / flagged words ----
print()
print('===== flagged words =====')
for kw in ['few-shot', 'significantly', 'novel', 'Crucially', 'modal', 'Fundus', 'vs. ResNet']:
    hits = [(i, p.text[:80]) for i, p in enumerate(doc.paragraphs) if kw.lower() in p.text.lower()]
    print(f'{kw!r}: {hits}')

# ---- P79 heading exact text ----
print()
for i, p in enumerate(doc.paragraphs):
    if p.style.name == 'heading2':
        print(f'heading2 P{i}: {p.text!r}')
