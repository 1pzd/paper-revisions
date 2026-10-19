# -*- coding: utf-8 -*-
"""C5-C9 fixes, DOCX part 2 (txt already applied). 2026-10-15."""
import sys, io, os
from docx import Document

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "CLIP.docx")
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

doc = Document(DOCX)

def find_para(substr, start=0):
    for i in range(start, len(doc.paragraphs)):
        if substr in doc.paragraphs[i].text:
            return i
    raise SystemExit("DOCX MISS anchor: %r" % substr[:60])

def has_math(p):
    return "<m:oMath" in p._p.xml

def rep_p(p, old, new, tag):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            print("[docx] ok(run)", tag)
            return
    if old in p.text:
        if has_math(p):
            raise SystemExit("DOCX MATH-WRITE BLOCKED [%s]" % tag)
        runs = p.runs
        full = p.text.replace(old, new)
        runs[0].text = full
        for r in runs[1:]:
            r.text = ""
        print("[docx] ok(rebuild)", tag)
        return
    raise SystemExit("DOCX MISS [%s]: %r" % (tag, old[:70]))

# 1. abstract C6/C9
i = find_para("we design a Bidirectional Rank-Aware Loss")
rep_p(doc.paragraphs[i], "we design a Bidirectional Rank-Aware Loss",
      "we incorporate a Bidirectional Rank-Aware Loss", "abs-design")

# 2. contribution C6+C9 (keep (e.g., Grade 0 < Grade 5) tail)
i = find_para("We propose a Bidirectional Rank-Aware Constraint")
rep_p(doc.paragraphs[i],
      "We propose a Bidirectional Rank-Aware Constraint that redefines the classification loss. By enforcing unimodal probability distributions and penalizing distance-based errors, we ensure predictions respect the strict logical progression of disease severity (e.g., Grade 0 < Grade 5).",
      "We incorporate a Bidirectional Rank-Aware Loss into the classification objective. The loss penalizes the expected absolute distance between the predicted distribution and the ground-truth grade, placing it in the established family of expected-distance ordinal objectives [14,15]; its bidirectional character reflects the symmetry of the absolute distance, which charges under- and overestimation alike. The contribution is therefore combinatorial: this rank objective is coupled with language-aligned semantic anchors and a multi-granularity encoder so that predictions respect the strict logical progression of disease severity (e.g., Grade 0 < Grade 5).",
      "contrib2")

# 3. related work C5
i = find_para("yet ophthalmic adoption remains limited")
rep_p(doc.paragraphs[i],
      "yet ophthalmic adoption remains limited [19,20]. Standard contrastive alignment",
      "yet ophthalmic adoption remains limited [19,20]. In cataract assessment, deep grading systems built on clinical scales have been developed for slit-lamp and retro-illumination photographs [21], nuclear cataract diagnosis from slit-lamp video [22], and cortical cataract grading with clinical sub-scores [23], but these pipelines remain CNN- or Transformer-based classifiers without language-aligned semantics. Standard contrastive alignment",
      "c5-clinical")
i = find_para("couples vision-language alignment with a bidirectional rank-aware objective")
rep_p(doc.paragraphs[i],
      "with a bidirectional rank-aware objective and grade-specific medical prompts",
      "with a Bidirectional Rank-Aware Loss and grade-specific medical prompts",
      "c5-tail")

# 4. framework component C9
i = find_para("Bidirectional Rank-Aware Optimization")
rep_p(doc.paragraphs[i],
      "Bidirectional Rank-Aware Optimization: To address the ordinal nature of cataracts, we introduce a specialized loss function that penalizes predictions based on their distance from the true grade, enforcing a unimodal probability distribution.",
      "Bidirectional Rank-Aware Loss: To address the ordinal nature of cataracts, we integrate a distance-weighted rank loss that penalizes predictions based on their expected ordinal distance from the true grade, encouraging a unimodal probability distribution.",
      "comp3-name")

# 5. fig captions C9
i = find_para("trained with the bidirectional rank-aware loss")
rep_p(doc.paragraphs[i], "trained with the bidirectional rank-aware loss",
      "trained with the Bidirectional Rank-Aware Loss", "c9-fig1")
i = find_para("of the rank-aware loss")
rep_p(doc.paragraphs[i], "of the rank-aware loss", "of the rank loss", "c9-fig2")

# 6. sec 3.4 name def C9
i = find_para("The bidirectional rank-aware term penalizes")
rep_p(doc.paragraphs[i], "The bidirectional rank-aware term penalizes",
      "The Bidirectional Rank-Aware Loss (hereafter the rank loss) penalizes", "34-name")

# 7. P069 append C6 (family + symmetry + combination)
i = find_para("Minimizing the rank term therefore concentrates")
doc.paragraphs[i].add_run(" This construction lies in the established family of expected-distance ordinal objectives [14,15]; the term “bidirectional” refers to the symmetry of the absolute distance rather than to a new loss family, and the contribution lies in coupling this rank objective with language-aligned semantic anchors and grade-specific medical prompts.")
print("[docx] ok append c6")

# 8. P071 (math) run-local C9 + append C8
i = find_para("This scheme utilizes Focal Loss")
rep_p(doc.paragraphs[i], "the rank-aware term", "the rank loss", "34-rankterm")
doc.paragraphs[i].add_run(" In the ablation study, the Imbalance Strategy switch jointly enables class-weighted sampling and the focal term; rows without it train with uniform sampling and the standard cross-entropy term in place of the focal term.")
print("[docx] ok append c8")

# 9. C7 baseline
i = find_para("Baseline (Row 1)")
rep_p(doc.paragraphs[i],
      "Baseline (Row 1): The naive application of CLIP yields an accuracy of 80.15%.",
      "Baseline (Row 1): The CLIP Base configuration—the global CLIP ViT-L/14 branch with grade-specific medical prompts, fine-tuned as described in Implementation Details with the local branch, rank loss, and imbalance handling all disabled—yields an accuracy of 80.15%, distinct from the zero-shot (64.20%) and linear-probe (79.50%) transfer settings in Table 2.",
      "c7-baseline")

# 10. C9 impact
i = find_para("Impact of Rank Loss")
rep_p(doc.paragraphs[i],
      "Impact of Rank Loss (Row 3): The integration of Rank Loss provides",
      "Impact of the Rank Loss (Row 3): The integration of the rank loss provides",
      "c9-impact")

# 11. C7/C8 table 3 caption
i = find_para("Ablation study demonstrating the incremental improvements")
doc.paragraphs[i].add_run(" Rank Loss denotes the Bidirectional Rank-Aware Loss; Imbalance Strategy denotes the joint use of class-weighted sampling and the focal loss term (rows without it train with uniform sampling and standard cross-entropy); CLIP Base denotes the fine-tuned global CLIP ViT-L/14 branch with medical prompts, distinct from the zero-shot and linear-probe rows of Table 2.")
print("[docx] ok append caption")

# 12. renumber [21]->[24], [22]->[25], [23]->[26]
i = find_para("gradient[21]")
rep_p(doc.paragraphs[i], "[21]", "[24]", "ref21")
i = find_para(")[22]")
rep_p(doc.paragraphs[i], "[22]", "[25]", "ref22")
i = find_para("augmentation[23]")
rep_p(doc.paragraphs[i], "[23]", "[26]", "ref23-aug")
rep_p(doc.paragraphs[i], "with a bidirectional rank-aware loss",
      "with the Bidirectional Rank-Aware Loss", "c9-conclusion")

# 13. insert 3 reference items before Lin focal loss
i = find_para("Focal loss for dense object detection")
anchor = doc.paragraphs[i]
new_refs = [
    "Son K Y ,Ko J ,Kim E , et al. Deep Learning-Based Cataract Detection and Grading from Slit-Lamp and Retro-Illumination Photographs[J].Ophthalmology Science,2022. DOI:10.1016/J.XOPS.2022.100147",
    "Shimizu E ,Tanji M ,Nakayama S , et al. AI-based diagnosis of nuclear cataract from slit-lamp videos[J].Scientific Reports,2023. DOI:10.1038/S41598-023-49563-7",
    "Wang J ,Xu Z ,Zheng W , et al. A Transformer-Based Knowledge Distillation Network for Cortical Cataract Grading[J].IEEE Transactions on Medical Imaging,2023. DOI:10.1109/TMI.2023.3327274",
]
for t in new_refs:
    anchor.insert_paragraph_before(t, style=anchor.style)
print("[docx] ok insert-3-refs")

doc.save(DOCX)
print("[docx] saved")

# verify
doc2 = Document(DOCX)
n_para = len(doc2.paragraphs)
n_tbl = len(doc2.tables)
n_math = sum(p._p.xml.count("<m:oMath") for p in doc2.paragraphs)
n_ref = sum(1 for p in doc2.paragraphs if p.style.name == "referenceitem")
print("[verify] paras=%d tables=%d oMath=%d refitems=%d" % (n_para, n_tbl, n_math, n_ref))
alltext = "\n".join(p.text for p in doc2.paragraphs)
for k in ["Son K Y", "Shimizu", "Wang J", "expected absolute distance",
          "hereafter the rank loss", "[24]", "[25]", "[26]", "CLIP Base configuration",
          "Bidirectional Rank-Aware Constraint", "Bidirectional Rank-Aware Optimization",
          "rank-aware term", "the rank-aware loss", "Naive application"]:
    print("[verify] %-38s %s" % (k, "FOUND(!!)" if k in alltext else "absent(ok)" if k in ("Bidirectional Rank-Aware Constraint","Bidirectional Rank-Aware Optimization","rank-aware term","the rank-aware loss","Naive application") else "ABSENT(!!)"))
print("DONE")
